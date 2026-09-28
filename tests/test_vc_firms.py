"""Insights > VC firms tests. Offline: the network, Firecrawl and the model are stubbed.

Protected: every firm has its own website read or a recorded reason why not;
post links are picked out of a news page (html, sitemap, Firecrawl markdown and
map) without nav, social or duplicate links; a post is stored once per firm and
sorted once; a skipped post stays pending; a firm with a site and a search gets
both reads; Firecrawl reads are made at most every `vc_firecrawl_hours` and
never with --free-only; the API hides Other unless asked, merges in funding
news that names the firm, prefers the firm's own post for one deal, filters by
where a post was found, and the Excel file has the same rows.
"""

from __future__ import annotations

import io
import re
from datetime import timedelta

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import delete, select

from app.config import settings
from app.db import session_scope
from app.ingest import vc_firms as vf
from app.llm.bedrock import LlmResult
from app.llm.budget import BudgetExceeded
from app.main import app
from app.models import (
    Article, Category, FundingRound, Stage, VcKind, VcPost, VcRead, utcnow,
)
from app.search import firecrawl as fc


@pytest.fixture(autouse=True)
def _clean():
    yield
    with session_scope() as s:
        s.execute(delete(VcPost))
        s.execute(delete(VcRead))


def firm(key: str) -> vf.VcFirm:
    found = vf.get_firm(key)
    assert found is not None, key
    return found


# --- Registry ---------------------------------------------------------------------

def test_registry_is_valid() -> None:
    firms = vf.load_firms()
    keys = [f.key for f in firms]
    assert len(firms) == 50
    assert len(keys) == len(set(keys))
    for f in firms:
        assert f.aliases, f.key
        for alias in f.aliases:
            re.compile(alias)
        # Every firm's own website is read, or the registry says why not.
        assert f.site or f.no_site, f.key
        assert f.reads, f.key
        if f.site:
            assert f.site.via in vf.SITE_VIAS and f.site.url.startswith("https://"), f.key
            if f.site.link_pattern:
                re.compile(f.site.link_pattern)
        if f.search:
            assert f.search.query, f.key
    assert sum(1 for f in firms if f.site) == 38


def test_bad_registry_entries_are_refused() -> None:
    with pytest.raises(ValueError, match="no_site"):
        vf._parse({"key": "x", "home": "https://x.test/", "search": "x"})
    with pytest.raises(ValueError, match="unknown site via"):
        vf._parse({"key": "x", "home": "https://x.test/",
                   "site": {"via": "search", "url": "https://x.test/"}})


def test_aliases_match_investor_lists() -> None:
    assert firm("peak-xv").matches("Sequoia India; Blume Ventures")
    assert firm("peak-xv").matches("Peak XV Partners")
    assert firm("z47").matches("Matrix Partners India")
    assert not firm("accel").matches("Accelerate Ventures")
    assert not firm("a91").matches(None)


# --- Reading a page ---------------------------------------------------------------

PAGE = """
<html><body>
<nav><a href="/about">About the firm and the team we have</a>
<a href="https://www.linkedin.com/company/x">Follow us on LinkedIn for more updates</a></nav>
<a href="/news/qzorbit-seed">Portfolio News Our Seed Investment in Qzorbit: payments for
  small shops September 24, 2026</a>
<a href="/news/qzorbit-seed">Read more about Our Seed Investment in Qzorbit: payments for small shops and more</a>
<a href="/news/short">Too short</a>
<a href="https://press.test/zyro">Zyro raises Rs 8 crore led by the firm today</a>
<script>var a = '<a href="/news/fake">Script link that is not a post at all</a>'</script>
</body></html>
"""


def test_posts_from_html_picks_post_links() -> None:
    found = vf.posts_from_html(PAGE, "https://vc.test/news")
    urls = [f.url for f in found]
    assert urls == ["https://vc.test/news/qzorbit-seed", "https://press.test/zyro"]
    qz = found[0]
    assert "Read more" not in qz.title and "2026" not in qz.title
    assert qz.published_at is not None and qz.date_approx     # a day from the card


def test_link_pattern_limits_to_posts() -> None:
    found = vf.posts_from_html(PAGE, "https://vc.test/news", pattern=r"vc\.test/news/.")
    assert [f.url for f in found] == ["https://vc.test/news/qzorbit-seed"]


def test_scraped_card_titles_are_tidied() -> None:
    # Lightspeed's cards arrive as "September 18\\ \\ Title" from Firecrawl.
    assert vf._tidy_title("September 18\\ \\ Collective Model Intelligence: next axis") == \
        "Collective Model Intelligence: next axis"
    # A title that merely starts with a month and a number is left alone.
    assert vf._tidy_title("May 5 things to know about Indian fintech") == \
        "May 5 things to know about Indian fintech"


def test_slug_title() -> None:
    url = "https://a.test/post/agrizy-raises-9-8m-series-a/"
    assert vf.slug_title(url) == "Agrizy raises 9 8m series a"
    assert vf.slug_title("https://a.test/blog-single.html") == "Blog single"


SITEMAP = """<?xml version="1.0"?><urlset>
<url><loc>https://vc.test/about</loc><lastmod>2026-09-20</lastmod></url>
<url><loc>https://vc.test/post/old-news-item</loc><lastmod>2026-01-02</lastmod></url>
<url><loc>https://vc.test/post/nimbus-raises-seed</loc><lastmod>2026-09-21T10:00:00Z</lastmod></url>
</urlset>"""


def test_posts_from_sitemap_newest_first() -> None:
    found, children = vf.posts_from_sitemap(SITEMAP, r"/post/.")
    assert children == []
    assert [f.url for f in found] == ["https://vc.test/post/nimbus-raises-seed",
                                      "https://vc.test/post/old-news-item"]
    assert found[0].title == "Nimbus raises seed" and found[0].date_approx


def test_sitemap_index_is_followed() -> None:
    index = ("<sitemapindex><sitemap><loc>https://vc.test/post-sitemap.xml</loc>"
             "</sitemap></sitemapindex>")
    read = vf.Read(via="sitemap", url="https://vc.test/sitemap.xml", link_pattern="/post/.")
    client = FakeClient({read.url: index, "https://vc.test/post-sitemap.xml": SITEMAP})
    assert len(vf._read_sitemap(read, client)) == 2


def test_map_read_filters_to_posts(monkeypatch) -> None:
    links = [fc.SiteLink("https://vc.test/", "Home", ""),
             fc.SiteLink("https://vc.test/news-and-announcement/gallabox-raises-12m",
                         "Gallabox raises $12M", "Conversational commerce"),
             fc.SiteLink("https://vc.test/content/blog/why-we-invested-in-klaar", "", "")]
    monkeypatch.setattr(vf.firecrawl, "map_site", lambda url: (links, 1, False))
    read = vf.Read(via="map", url="https://vc.test",
                   link_pattern="/(news-and-announcement|content/blog)/.")
    found, credits = vf._read_map(read)
    assert credits == 1
    assert [f.title for f in found] == ["Gallabox raises $12M", "Why we invested in klaar"]


def test_posts_from_markdown() -> None:
    md = ("[![logo](https://vc.test/logo.png)](https://vc.test/)\n"
          "- [Nimbus raises $3M seed round to build farm robots](https://vc.test/blog/nimbus)\n"
          "- [Home](https://vc.test/)\n"
          "- [**Fund IV** announced for deep tech founders in India](https://vc.test/blog/fund-iv \"t\")")
    found = vf.posts_from_markdown(md, "https://vc.test/blog")
    assert [f.url for f in found] == ["https://vc.test/blog/nimbus", "https://vc.test/blog/fund-iv"]
    assert found[1].title.startswith("Fund IV")


# --- Reading firms ---------------------------------------------------------------

class FakeResponse:
    def __init__(self, text, url):
        self.text, self.content, self.url = text, text.encode(), url
        self.status_code = 200

    def raise_for_status(self):
        pass


class FakeClient:
    def __init__(self, pages):
        self.pages = pages

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def get(self, url):
        return FakeResponse(self.pages[url], url)


def test_html_read_stores_once_per_firm(monkeypatch) -> None:
    accel = firm("accel")
    page = PAGE.replace("/news/", "https://www.accel.com/news/")
    monkeypatch.setattr(vf, "_client", lambda: FakeClient({accel.site.url: page}))
    monkeypatch.setattr(vf, "_fill_from_pages", lambda ids, client: 0)

    first = vf.read_firm(accel)
    again = vf.read_firm(accel)
    assert first.error is None and first.items_new == 1
    assert again.items_new == 0
    with session_scope() as s:
        assert s.scalar(select(VcPost.via).where(VcPost.firm == "accel")) == "html"


def test_failed_read_is_recorded_not_raised(monkeypatch) -> None:
    accel = firm("accel")
    monkeypatch.setattr(vf, "_client", lambda: FakeClient({accel.site.url: "<html></html>"}))
    run = vf.read_firm(accel)
    assert run.error and "no posts found" in run.error


def stub_firecrawl(monkeypatch):
    calls = []

    def fake_search(query, limit=None, kind="news", tbs=None, steered=True):
        calls.append(("search", query, tbs))
        return fc.Search(query=query, searched=query, kind=kind, credits_used=2, results=[
            fc.SearchResult(n=1, url=f"https://n.test/{len(calls)}",
                            title="Firm backs Nimbus in $3M seed round",
                            description="", domain="n.test", published="2 days ago")])

    def fake_scrape(url):
        calls.append(("scrape", url, None))
        md = f"[Nimbus raises $3M seed round to build farm robots]({url.rstrip('/')}/nimbus)"
        return fc.Page(url=url, title="", description="", markdown=md, credits_used=1), False

    monkeypatch.setattr(vf.firecrawl, "search", fake_search)
    monkeypatch.setattr(vf.firecrawl, "scrape", fake_scrape)
    return calls


def test_firecrawl_firms_are_spaced_out(monkeypatch) -> None:
    calls = stub_firecrawl(monkeypatch)

    first = vf.refresh_firms(only="tiger-global", classify=False)
    assert first.credits == 2 and first.new == 1
    assert calls[0][2] == "qdr:m"                 # first search looks back a month

    second = vf.refresh_firms(only="tiger-global", classify=False)
    assert second.reads == [] and second.skipped == 1
    assert len(calls) == 1                         # no second paid search

    forced = vf.refresh_firms(only="tiger-global", force=True, classify=False)
    assert len(forced.reads) == 1 and calls[-1][2] == "qdr:w"


def test_firecrawl_due_after_interval(monkeypatch) -> None:
    stub_firecrawl(monkeypatch)
    vf.refresh_firms(only="omnivore", classify=False)
    later = utcnow() + timedelta(hours=settings.vc_firecrawl_hours)
    omnivore = firm("omnivore")
    assert not vf.firecrawl_due(omnivore, omnivore.site)
    assert vf.firecrawl_due(omnivore, omnivore.site, now=later)


def test_firm_with_site_and_search_gets_both_reads(monkeypatch) -> None:
    calls = stub_firecrawl(monkeypatch)        # Lightspeed: scrape + search
    result = vf.refresh_firms(only="lightspeed-india", classify=False)
    assert sorted(r.via for r in result.reads) == ["scrape", "search"]
    assert sorted(c[0] for c in calls) == ["scrape", "search"]
    with session_scope() as s:
        vias = set(s.scalars(select(VcPost.via).where(VcPost.firm == "lightspeed-india")))
    assert vias == {"scrape", "search"}


def test_free_only_still_reads_free_sites(monkeypatch) -> None:
    calls = stub_firecrawl(monkeypatch)        # Elevation: free html + paid search
    elevation = firm("elevation")
    page = ('<a href="https://www.elevationcapital.com/perspectives/investing-in-plazza">'
            'Investing in Plazza, an on-demand pharmacy</a>')
    monkeypatch.setattr(vf, "_client", lambda: FakeClient({elevation.site.url: page}))
    monkeypatch.setattr(vf, "_fill_from_pages", lambda ids, client: 0)
    result = vf.refresh_firms(only="elevation", free_only=True, classify=False)
    assert [r.via for r in result.reads] == ["html"] and result.skipped == 1
    assert calls == []


def test_free_only_never_uses_firecrawl(monkeypatch) -> None:
    calls = stub_firecrawl(monkeypatch)
    result = vf.refresh_firms(only="endiya", free_only=True, classify=False)
    assert result.reads == [] and result.skipped == 1 and calls == []


def test_failed_firecrawl_read_is_retried(monkeypatch) -> None:
    def broken(*a, **kw):
        raise fc.FirecrawlError("Firecrawl credits are used up. Top up at firecrawl.dev.")

    monkeypatch.setattr(vf.firecrawl, "search", broken)
    run = vf.refresh_firms(only="a91", classify=False).reads[0]
    assert "credits are used up" in run.error
    a91 = firm("a91")
    assert vf.firecrawl_due(a91, a91.search)      # a failure does not count as a read


# --- Classifying ----------------------------------------------------------------

def add_posts(*specs) -> list[int]:
    ids = []
    with session_scope() as s:
        for firm_key, title in specs:
            p = VcPost(firm=firm_key, url=f"https://p.test/{len(ids)}-{firm_key}",
                       title=title, via="html", published_at=utcnow() - timedelta(days=1))
            s.add(p)
            s.flush()
            ids.append(p.id)
    return ids


def stub_model(monkeypatch, answers: dict[str, dict]):
    prompts = []

    def fake(feature, prompt, **kw):
        prompts.append(prompt)
        out = []
        for line in prompt.splitlines():
            m = re.match(r"\[(\d+)\] \(Firm: [^)]+\) (\S+)", line)
            if m and m.group(2) in answers:
                out.append({"i": int(m.group(1)), **answers[m.group(2)]})
        return out, LlmResult("", settings.model_cheap, 900, 150, 0.0002)

    monkeypatch.setattr(vf, "invoke_json", fake)
    return prompts


def test_classify_sorts_once_and_keeps_skipped_pending(monkeypatch) -> None:
    add_posts(("accel", "Nimbus raises $3M seed"), ("accel", "Podcast: on patience"),
              ("accel", "Garbled item"))
    prompts = stub_model(monkeypatch, {
        "Nimbus": {"kind": "Investment", "company": "Nimbus", "round": "seed",
                   "amount": "$3M", "headline": "Nimbus raises $3M seed round"},
        "Podcast:": {"kind": "podcast", "company": "null", "headline": "On patience"},
    })
    result = vf.classify_pending()
    assert result.classified == 2 and result.considered == 3
    with session_scope() as s:
        posts = {p.title.split()[0]: p for p in s.scalars(select(VcPost)).all()}
    assert posts["Nimbus"].kind == VcKind.INVESTMENT and posts["Nimbus"].round_label == "Seed"
    assert posts["Podcast:"].kind == VcKind.OTHER and posts["Podcast:"].company is None
    assert posts["Garbled"].classified_at is None          # retried next time

    vf.classify_pending()
    assert "Nimbus" not in prompts[-1]                      # never paid for twice


def test_classify_stops_cleanly_on_budget(monkeypatch) -> None:
    add_posts(("accel", "Nimbus raises $3M seed"))

    def capped(*a, **kw):
        raise BudgetExceeded("daily cap reached", None)

    monkeypatch.setattr(vf, "invoke_json", capped)
    result = vf.classify_pending()
    assert result.classified == 0 and "cap" in result.stopped_reason


# --- API ---------------------------------------------------------------------------

@pytest.fixture
def news_round():
    """A Funding story from the feed whose investors name Blume and Accel."""
    with session_scope() as s:
        a = Article(url="https://feed.test/orbo", url_hash="vcorbo", source="entrackr",
                    headline="Orbo raises $2M seed from Blume Ventures and Accel",
                    category=Category.FUNDING, published_at=utcnow() - timedelta(days=2))
        s.add(a)
        s.flush()
        s.add(FundingRound(article_id=a.id, company="Orbo", stage=Stage.SEED,
                           round_label="Seed", amount="$2M",
                           investors="Blume Ventures; Accel"))
        article_id = a.id
    yield article_id
    with session_scope() as s:
        s.execute(delete(FundingRound).where(FundingRound.article_id == article_id))
        s.execute(delete(Article).where(Article.id == article_id))


def sorted_post(firm_key, kind, company, days_ago=1, headline=None, via="html"):
    with session_scope() as s:
        s.add(VcPost(firm=firm_key,
                     url=f"https://site.test/{firm_key}/{company}/{kind.name}/{via}",
                     title=headline or f"{company} post", via=via, kind=kind,
                     headline=headline or f"{company} post", company=company,
                     published_at=utcnow() - timedelta(days=days_ago),
                     classified_at=utcnow()))


def test_api_hides_other_and_merges_news(news_round) -> None:
    sorted_post("accel", VcKind.INVESTMENT, "Orbo", headline="Our seed investment in Orbo")
    sorted_post("accel", VcKind.OTHER, None, headline="Podcast on patience")
    client = TestClient(app)

    body = client.get("/api/insights/vcs").json()
    assert body["hidden_other"] == 1
    by_firm = {}
    for p in body["posts"]:
        by_firm.setdefault(p["firm"], []).append(p)
    # Accel's own post covers the Orbo round, so the news copy is dropped.
    assert [p["origin"] for p in by_firm["accel"]] == ["site"]
    # Blume has no post of its own; the feed's round shows under it.
    assert by_firm["blume"][0]["origin"] == "news"
    assert by_firm["blume"][0]["source_label"] == "Entrackr"
    counts = {f["key"]: f["count"] for f in body["firms"]}
    assert counts["accel"] == 1 and counts["blume"] == 1

    everything = client.get("/api/insights/vcs", params={"all": "true", "firm": "accel"}).json()
    assert len(everything["posts"]) == 2
    assert {p["firm"] for p in everything["posts"]} == {"accel"}


def test_api_prefers_site_and_filters_by_origin() -> None:
    sorted_post("lightspeed-india", VcKind.INVESTMENT, "Nimbus", via="search",
                headline="Lightspeed backs Nimbus, reports say")
    sorted_post("lightspeed-india", VcKind.INVESTMENT, "Nimbus", via="scrape",
                headline="Our investment in Nimbus")
    client = TestClient(app)
    rows = client.get("/api/insights/vcs", params={"firm": "lightspeed-india"}).json()["posts"]
    assert [(r["origin"], r["source_label"]) for r in rows] == [("site", "Firm website")]

    site_only = client.get("/api/insights/vcs", params={"origin": "site"}).json()
    assert {p["origin"] for p in site_only["posts"]} == {"site"}
    assert client.get("/api/insights/vcs", params={"origin": "blog"}).status_code == 400

    firms = {f["key"]: f for f in site_only["firms"]}
    assert [r["via"] for r in firms["lightspeed-india"]["reads"]] == ["scrape", "search"]
    assert firms["tiger-global"]["no_site"]


def test_api_date_window_and_unknown_firm(news_round) -> None:
    sorted_post("accel", VcKind.FUND, "Fund IV", days_ago=40)
    client = TestClient(app)
    today = utcnow().date()
    recent = client.get("/api/insights/vcs", params={
        "firm": "accel", "start": str(today - timedelta(days=10))}).json()
    assert all(p["company"] != "Fund IV" for p in recent["posts"])
    assert client.get("/api/insights/vcs", params={"firm": "nope"}).status_code == 404


def test_excel_has_the_same_rows(news_round) -> None:
    from openpyxl import load_workbook

    sorted_post("accel", VcKind.PORTFOLIO, "Rentomojo", headline="Rentomojo lists on NSE")
    client = TestClient(app)
    screen = client.get("/api/insights/vcs").json()["posts"]
    resp = client.get("/api/insights/vcs.xlsx")
    assert resp.status_code == 200
    ws = load_workbook(io.BytesIO(resp.content)).active
    assert ws.max_row - 1 == len(screen)
    assert ws.cell(row=1, column=1).value == "VC firm"
