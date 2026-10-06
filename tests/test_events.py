"""Insights > Events organised tests. Offline: Apify, Firecrawl and the model are stubbed.

Protected: firms' Luma calendars come from the registry; Apify is called with
the calendar, a period and both caps, and its errors read as sentences; Luma
rows map to events, with diagnostic rows never stored and global calendars
kept to India; a pasted page is read by the model once per change of its text,
and the model's answer is untrusted (no date, no event; only http(s) links);
a re-read updates an event instead of adding it twice; upcoming reads are made
at most every `events_hours` and past ones once per calendar; --skip-paid reads
nothing; the API splits upcoming from past, filters by organiser, exports the
same rows to Excel, and adds and removes pasted sources; "Refresh now" has an
events scope.
"""

from __future__ import annotations

import io
import time
from datetime import datetime, timedelta, timezone

import httpx
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import delete, func, select

from app import scheduler
from app.config import settings
from app.db import session_scope
from app.ingest import events as ev
from app.ingest import vc_firms as vf
from app.llm.bedrock import LlmResult
from app.main import app
from app.models import Event, EventRead, EventSource, utcnow
from app.search import apify
from app.search import firecrawl as fc

IST = timezone(timedelta(hours=5, minutes=30))

# The actor's own documented example row, trimmed.
LUMA_ROW = {
    "eventId": "evt-DglQU0ckykBp1nQ", "url": "https://lu.ma/ou0z635y",
    "name": "Surge Open House: Bengaluru",
    "startAt": "2026-10-20T12:30:00.000Z", "endAt": "2026-10-20T15:30:00.000Z",
    "timezone": "Asia/Kolkata", "isOnline": False, "venueName": "Peak XV office",
    "addressShort": "Indiranagar, Bengaluru", "city": "Bengaluru", "country": "India",
    "countryCode": "IN", "hostName": "Peak XV", "calendarName": "Peak XV",
    "description": "Inviting founders in Bengaluru to discover Surge.",
}


@pytest.fixture(autouse=True)
def _clean():
    yield
    with session_scope() as s:
        s.execute(delete(Event))
        s.execute(delete(EventRead))
        s.execute(delete(EventSource))


@pytest.fixture
def token(monkeypatch):
    monkeypatch.setattr(settings, "apify_api_token", "apify-test")


def feed_for(firm: str) -> ev.Feed:
    return next(f for f in ev.registry_feeds() if f.firm == firm)


def stub_actor(monkeypatch, rows_by_period: dict[str, list[dict]]):
    calls = []

    def fake(actor, run_input, *, max_items, max_charge_usd=None):
        calls.append({"actor": actor, "input": run_input, "max_items": max_items,
                      "max_charge_usd": max_charge_usd})
        return list(rows_by_period.get(run_input["period"], []))

    monkeypatch.setattr(ev.apify, "run_actor", fake)
    return calls


def stub_page(monkeypatch, markdown: str, answer: list[dict]):
    """Firecrawl returns `markdown`; the model answers `answer`. Returns the
    list of prompts the model was sent."""
    prompts = []

    def scrape(url):
        return fc.Page(url=url, title="Events", description="", markdown=markdown,
                       credits_used=1), False

    def model(feature, prompt, **kwargs):
        prompts.append(prompt)
        return answer, LlmResult(text="", model="m", tokens_in=1, tokens_out=1, cost_usd=0.0009)

    monkeypatch.setattr(ev.firecrawl, "scrape", scrape)
    monkeypatch.setattr(ev, "invoke_json", model)
    return prompts


def add_event(feed: ev.Feed, title: str, starts: datetime, **kw) -> None:
    found = ev.Found(url=kw.pop("url", f"https://luma.com/{title.lower().replace(' ', '-')}"),
                     title=title, starts_at=starts, **kw)
    ev.store(feed, [found], "apify")


# --- Registry ---------------------------------------------------------------------

def test_registry_has_event_sources_for_firms() -> None:
    feeds = ev.registry_feeds()
    firms = {f.firm for f in feeds}
    assert {"peak-xv", "accel", "general-catalyst-india", "elevation", "kae",
            "together-fund", "blume", "z47", "fireside", "venture-catalysts"} <= firms
    # Apify only gets calendars it can read; one with only a /calendar/cal-...
    # address is read as a page.
    for f in feeds:
        if f.via == "luma":
            assert ev.actor_can_read(f.url), f.url
        elif ev.is_luma(f.url):
            assert f.via == "page", f.url
    assert {f.via for f in feeds} == {"luma", "page", "html"}
    assert len({f.key for f in feeds}) == len(feeds)
    # Every firm with a live site has its home page as a source; the rest say why.
    for firm in vf.load_firms():
        if firm.no_events:
            assert firm.key not in firms, firm.key
        else:
            assert firm.key in firms, firm.key
    homes = [f for f in feeds if f.site]
    assert all(f.via in ("html", "page") for f in homes)
    assert len(homes) >= 40
    assert next(f for f in homes if f.firm == "accel").india_only
    # Global calendars are kept to India; an India-first one is not filtered.
    assert feed_for("general-catalyst-india").india_only
    assert not feed_for("peak-xv").india_only
    # Keys are stable across loads, so read history survives a restart.
    assert [f.key for f in ev.registry_feeds()] == [f.key for f in feeds]


def test_registry_rejects_a_bad_events_entry() -> None:
    raw = {"key": "x", "home": "https://x.test", "no_site": "none",
           "events": [{"via": "ics", "url": "https://x.test/cal"}]}
    with pytest.raises(ValueError):
        vf._parse(raw)


# --- Apify client -------------------------------------------------------------------

def _apify_with(monkeypatch, handler):
    real = httpx.Client
    monkeypatch.setattr(apify, "_client", lambda: real(transport=httpx.MockTransport(handler)))


def test_apify_runs_the_actor_with_both_caps(monkeypatch, token) -> None:
    seen = {}

    def handler(request: httpx.Request) -> httpx.Response:
        seen["path"] = request.url.path
        seen["params"] = dict(request.url.params)
        seen["auth"] = request.headers.get("authorization")
        seen["body"] = request.content
        return httpx.Response(201, json=[LUMA_ROW, "junk"])

    _apify_with(monkeypatch, handler)
    items = apify.run_actor("dami_studio~luma-events-scraper", {"startUrls": ["u"]},
                            max_items=25, max_charge_usd=0.05)
    assert items == [LUMA_ROW]
    assert seen["path"] == "/v2/actors/dami_studio~luma-events-scraper/run-sync-get-dataset-items"
    assert seen["params"]["maxItems"] == "25"
    assert seen["params"]["maxTotalChargeUsd"] == "0.0500"
    assert seen["auth"] == "Bearer apify-test"
    assert b'"startUrls"' in seen["body"]


@pytest.mark.parametrize("status, words", [
    (401, "rejected the token"), (402, "credits or plan limit"),
    (408, "did not finish"), (429, "rate limit"), (500, "could not do that"),
])
def test_apify_errors_read_as_sentences(monkeypatch, token, status, words) -> None:
    _apify_with(monkeypatch, lambda request: httpx.Response(status, json={}))
    with pytest.raises(apify.ApifyError, match=words):
        apify.run_actor("a~b", {}, max_items=1)


def test_apify_without_a_token_makes_no_request(monkeypatch) -> None:
    def handler(request):  # pragma: no cover - must not be reached
        raise AssertionError("a request was made")

    _apify_with(monkeypatch, handler)
    with pytest.raises(apify.ApifyError, match="not set up"):
        apify.run_actor("a~b", {}, max_items=1)


# --- Luma rows ----------------------------------------------------------------------

def test_luma_row_becomes_an_event() -> None:
    found = ev.from_luma(LUMA_ROW)
    assert found.title == "Surge Open House: Bengaluru"
    assert found.starts_at == datetime(2026, 10, 20, 12, 30, tzinfo=timezone.utc)
    assert found.ends_at == datetime(2026, 10, 20, 15, 30, tzinfo=timezone.utc)
    assert (found.city, found.venue, found.online) == ("Bengaluru", "Peak XV office", False)
    assert found.host == "Peak XV" and not found.date_only


def test_diagnostic_and_sample_rows_are_not_events() -> None:
    assert ev.from_luma({"_diagnostic": True, "errorCode": "NOT_FOUND"}) is None
    assert ev.from_luma({**LUMA_ROW, "_sample": True}) is None
    assert ev.from_luma({**LUMA_ROW, "url": "javascript:alert(1)"}) is None


def test_read_luma_asks_for_the_period_within_caps(monkeypatch) -> None:
    calls = stub_actor(monkeypatch, {"future": [LUMA_ROW], "past": []})
    run = EventRead(apify_usd=0.0)
    found = ev.read_luma(feed_for("peak-xv"), "future", run)

    assert [f.title for f in found] == ["Surge Open House: Bengaluru"]
    call = calls[0]
    assert call["actor"] == settings.apify_luma_actor
    # The actor rejects luma.com links (BAD_URL); it is sent the lu.ma one.
    assert call["input"]["startUrls"] == ["https://lu.ma/peakxvpartners"]
    assert call["input"]["period"] == "future"
    assert call["max_items"] == settings.events_upcoming_items
    assert call["max_charge_usd"] <= settings.apify_max_charge_usd
    assert run.apify_usd == pytest.approx(settings.apify_price_per_event
                                          + settings.apify_price_per_run)


def test_global_calendar_keeps_india_only(monkeypatch) -> None:
    abroad = {**LUMA_ROW, "url": "https://lu.ma/sf1", "name": "SF dinner",
              "countryCode": "US", "timezone": "America/Los_Angeles"}
    stub_actor(monkeypatch, {"future": [LUMA_ROW, abroad]})
    found = ev.read_luma(feed_for("general-catalyst-india"), "future", EventRead())
    assert [f.title for f in found] == ["Surge Open House: Bengaluru"]


def test_a_profile_link_is_explained(monkeypatch) -> None:
    stub_actor(monkeypatch, {"future": [{"_diagnostic": True, "errorCode": "UNSUPPORTED_URL"}]})
    with pytest.raises(ValueError, match="personal profile"):
        ev.read_luma(feed_for("peak-xv"), "future", EventRead())


def test_links_are_sent_as_the_actor_accepts_them() -> None:
    assert ev.actor_url("https://luma.com/AccelAtoms") == "https://lu.ma/AccelAtoms"
    assert ev.actor_url("https://www.luma.com/x?k=1") == "https://lu.ma/x?k=1"
    assert ev.actor_url("http://lu.ma/x") == "https://lu.ma/x"
    # A calendar with no short name cannot be read by the actor at all.
    assert ev.actor_can_read("https://luma.com/generalcatalyst")
    assert not ev.actor_can_read("https://luma.com/calendar/cal-8RP6LcGzY0D7R8R")
    assert not ev.actor_can_read("https://firm.test/events")


def test_a_pasted_calendar_with_no_short_name_is_read_as_a_page() -> None:
    feed = ev.add_source("https://luma.com/calendar/cal-8RP6LcGzY0D7R8R", label="Bessemer")
    assert feed.via == "page"
    assert ev.add_source("https://luma.com/founders-club").via == "luma"


def test_the_actors_own_words_are_kept(monkeypatch) -> None:
    stub_actor(monkeypatch, {"future": [{"_diagnostic": True, "errorCode": "BAD_URL",
                                         "error": "That link is not on lu.ma."}]})
    with pytest.raises(ValueError, match=r"BAD_URL: That link is not on lu\.ma\."):
        ev.read_luma(feed_for("peak-xv"), "future", EventRead())


def test_a_calendar_with_no_events_is_not_an_error(monkeypatch) -> None:
    stub_actor(monkeypatch, {"future": [{"_diagnostic": True, "errorCode": "NO_RESULTS"}]})
    assert ev.read_luma(feed_for("peak-xv"), "future", EventRead()) == []


# --- Dates and pages ----------------------------------------------------------------

def test_parse_when() -> None:
    # A time with no zone is Indian time.
    assert ev.parse_when("2026-10-14T18:30") == (
        datetime(2026, 10, 14, 13, 0, tzinfo=timezone.utc), False)
    # A bare day is flagged, never given a made-up hour.
    when, day_only = ev.parse_when("2026-10-14")
    assert day_only and when.astimezone(IST).date().isoformat() == "2026-10-14"
    assert ev.parse_when("2026-10-14T08:00:00Z")[0] == datetime(2026, 10, 14, 8, tzinfo=timezone.utc)
    assert ev.parse_when("14 Oct 2026")[1] is True
    assert ev.parse_when("soon") == (None, False)
    assert ev.parse_when(None) == (None, False)


def test_page_answer_is_untrusted() -> None:
    base = "https://firm.test/events"
    found = ev.coerce_page_events({"events": [
        {"title": "Demo Day", "start": "2026-11-02", "url": "/events/demo-day", "online": False},
        {"title": "No date", "start": None},
        {"title": "Sneaky", "start": "2026-11-03T10:00", "url": "javascript:alert(1)"},
        {"start": "2026-11-04"},
        "junk",
    ]}, base)
    assert [f.title for f in found] == ["Demo Day", "Sneaky"]
    assert found[0].url == "https://firm.test/events/demo-day" and found[0].own_link
    assert found[0].date_only and found[0].online is False
    assert found[1].url == base and not found[1].own_link


def test_a_page_is_read_by_the_model_once_per_change(monkeypatch) -> None:
    feed = ev.add_source("https://firm.test/events", firm="kalaari")
    answer = [{"title": "Konverse 2026", "start": "2026-11-20T18:30", "city": "Bengaluru"}]
    prompts = stub_page(monkeypatch, "## Konverse 2026, 20 Nov", answer)

    first = ev.read_feed(feed, "page")
    assert (first.error, first.items_new, first.credits_used) == (None, 1, 1)
    assert first.llm_usd == pytest.approx(0.0009)
    assert "Konverse 2026" in prompts[0] and "Kalaari" in prompts[0]

    again = ev.read_feed(feed, "page")           # same text: no model call
    assert len(prompts) == 1 and again.items_new == 0 and again.llm_usd == 0

    stub_page(monkeypatch, "## Konverse 2026, 21 Nov (moved)", answer)
    ev.read_feed(feed, "page")                   # changed text: read again
    with session_scope() as s:
        assert s.scalar(select(func.count(Event.id))) == 1


def test_a_home_page_without_events_costs_nothing(monkeypatch) -> None:
    feed = next(f for f in ev.registry_feeds() if f.site and f.via == "html")
    monkeypatch.setattr(ev, "fetch_page", lambda url: (
        url, "<main><h1>We back bold founders</h1><p>Portfolio, team and news.</p></main>"))
    monkeypatch.setattr(ev, "invoke_json", lambda *a, **k: pytest.fail("the model was called"))
    run = ev.read_feed(feed, "page")
    assert (run.error, run.items_seen, run.llm_usd, run.credits_used) == (None, 0, 0, 0)


def test_a_rate_limited_page_read_waits_and_retries_once(monkeypatch) -> None:
    feed = ev.add_source("https://firm.test/events", label="Firm")
    calls = []

    def scrape(url):
        calls.append(url)
        if len(calls) == 1:
            raise fc.FirecrawlError("Firecrawl's rate limit was hit. Wait a moment and try again.")
        return fc.Page(url=url, title="", description="", markdown="nothing", credits_used=1), False

    monkeypatch.setattr(ev.firecrawl, "scrape", scrape)
    monkeypatch.setattr(ev, "RATE_LIMIT_WAIT", 0)
    monkeypatch.setattr(ev, "invoke_json", lambda *a, **k: (
        [], LlmResult(text="", model="m", tokens_in=1, tokens_out=1, cost_usd=0.0)))
    run = ev.read_feed(feed, "page")
    assert run.error is None and len(calls) == 2


def test_mentions_events() -> None:
    assert ev.mentions_events("Founders Summit, 14 Nov 2026, Bengaluru")
    assert ev.mentions_events("Join our Demo Day on Oct 9")
    assert not ev.mentions_events("Founders Summit -- our thinking on growth")   # no date
    assert not ev.mentions_events("Raised Series A on 14 Nov 2026")           # no event


def test_home_pages_are_checked_weekly() -> None:
    home = next(f for f in ev.registry_feeds() if f.site)
    calendar = feed_for("peak-xv")
    assert home.every == timedelta(hours=settings.events_site_hours) > calendar.every
    with session_scope() as s:
        s.add(EventRead(source_key=home.key, kind="page",
                        started_at=utcnow() - timedelta(days=2)))
    assert ev.due_kinds(home) == []                     # read 2 days ago: not due
    assert ev.due_kinds(home, now=utcnow() + timedelta(days=6)) == ["page"]
    assert home.via_label.startswith("Firm website")


def test_a_global_firm_page_asks_for_india_only() -> None:
    accel = next(f for f in ev.registry_feeds() if f.site and f.firm == "accel")
    assert "Only list events held in India" in ev.page_prompt(accel, accel.url, "text")
    fireside = next(f for f in ev.registry_feeds() if f.firm == "fireside")
    assert "Only list events held in India" not in ev.page_prompt(fireside, fireside.url, "x")


def test_an_html_page_is_read_free_with_its_links(monkeypatch) -> None:
    feed = next(f for f in ev.registry_feeds() if f.via == "html" and not f.site)
    day = (utcnow() + timedelta(days=12)).astimezone(IST).date().isoformat()
    page = (f"<html><nav>Menu</nav><main><h2>Events</h2><p>{day}</p>"
            "<a href='/events/demo-day'>Demo Day</a><script>x()</script></main></html>")
    monkeypatch.setattr(ev, "fetch_page", lambda url: ("https://firm.test/events", page))
    monkeypatch.setattr(ev.firecrawl, "scrape", lambda url: pytest.fail("Firecrawl was called"))
    prompts = []

    def model(feature, prompt, **kwargs):
        prompts.append(prompt)
        return ([{"title": "Demo Day", "start": day, "url": "https://firm.test/events/demo-day"}],
                LlmResult(text="", model="m", tokens_in=1, tokens_out=1, cost_usd=0.0007))

    monkeypatch.setattr(ev, "invoke_json", model)
    run = ev.read_feed(feed, "page")

    assert (run.error, run.items_new, run.credits_used) == (None, 1, 0)
    # Links reach the model as [text](url); menus and scripts do not.
    assert "[Demo Day](https://firm.test/events/demo-day)" in prompts[0]
    assert "Menu" not in prompts[0] and "x()" not in prompts[0]
    events = TestClient(app).get("/api/insights/events").json()["events"]
    assert [(e["title"], e["via_label"]) for e in events] == [("Demo Day", "Website")]


def test_a_rescheduled_event_is_updated_not_added(monkeypatch) -> None:
    feed = feed_for("peak-xv")
    stub_actor(monkeypatch, {"future": [LUMA_ROW]})
    ev.read_feed(feed, "upcoming")
    moved = {**LUMA_ROW, "startAt": "2026-10-27T12:30:00.000Z", "endAt": None}
    stub_actor(monkeypatch, {"future": [moved]})
    run = ev.read_feed(feed, "upcoming")

    assert run.items_new == 0
    with session_scope() as s:
        rows = s.scalars(select(Event)).all()
        assert len(rows) == 1
        assert rows[0].starts_at == datetime(2026, 10, 27, 12, 30, tzinfo=timezone.utc)
        assert rows[0].firm == "peak-xv"


# --- When reads are due --------------------------------------------------------------

def test_new_calendar_reads_upcoming_and_backfills_past_once(monkeypatch) -> None:
    feed = feed_for("accel")
    assert ev.due_kinds(feed) == ["upcoming", "past"]

    stub_actor(monkeypatch, {"future": [LUMA_ROW], "past": []})
    for kind in ev.due_kinds(feed):
        ev.read_feed(feed, kind)
    assert ev.due_kinds(feed) == []

    later = utcnow() + timedelta(hours=settings.events_hours, minutes=1)
    assert ev.due_kinds(feed, now=later) == ["upcoming"]       # past is not re-read
    assert ev.due_kinds(feed, force=True) == ["upcoming", "past"]


def test_a_failed_read_is_retried(monkeypatch) -> None:
    feed = feed_for("accel")

    def broken(*args, **kwargs):
        raise apify.ApifyError("Apify credits or plan limit used up.")

    monkeypatch.setattr(ev.apify, "run_actor", broken)
    run = ev.read_feed(feed, "upcoming")
    assert run.error == "Apify credits or plan limit used up."
    assert "upcoming" in ev.due_kinds(feed)


def test_skip_paid_reads_nothing(monkeypatch) -> None:
    calls = stub_actor(monkeypatch, {"future": [LUMA_ROW]})
    result = ev.refresh_events(free_only=True)
    assert result.reads == [] and calls == []
    assert result.skipped == len(ev.all_feeds())


def test_refresh_reads_one_firm(monkeypatch) -> None:
    calls = stub_actor(monkeypatch, {"future": [LUMA_ROW], "past": []})
    result = ev.refresh_events(only="peak-xv")
    assert {c["input"]["startUrls"][0] for c in calls} == {"https://lu.ma/peakxvpartners"}
    assert result.new == 1 and result.failed == 0
    with pytest.raises(ValueError):
        ev.refresh_events(only="nobody")


# --- API ---------------------------------------------------------------------------

def test_api_splits_upcoming_and_past() -> None:
    now = utcnow()
    peak = feed_for("peak-xv")
    add_event(peak, "Later", now + timedelta(days=9))
    add_event(peak, "Soon", now + timedelta(days=2))
    add_event(peak, "Last month", now - timedelta(days=30))
    add_event(peak, "Last week", now - timedelta(days=7))
    # A day-only event today is still upcoming until the day is over.
    today = datetime.combine(now.astimezone(IST).date(), datetime.min.time(), IST)
    add_event(peak, "Today", today.astimezone(timezone.utc), date_only=True)

    client = TestClient(app)
    upcoming = client.get("/api/insights/events").json()
    assert [e["title"] for e in upcoming["events"]] == ["Today", "Soon", "Later"]
    assert upcoming["counts"] == {"upcoming": 3, "past": 2}
    assert upcoming["events"][0]["organiser"] == "Peak XV Partners"
    assert upcoming["events"][0]["via_label"] == "Luma, via Apify"

    past = client.get("/api/insights/events", params={"when": "past"}).json()
    assert [e["title"] for e in past["events"]] == ["Last week", "Last month"]

    # Every organiser with a source is listed, with its count for the tab.
    organisers = {o["key"]: o["count"] for o in upcoming["organisers"]}
    assert organisers["peak-xv"] == 3 and organisers["accel"] == 0
    # Every tracked firm is an organiser, with or without a source.
    assert set(organisers) >= {f.key for f in vf.load_firms()}
    firms = {f["key"]: f for f in upcoming["firms"]}
    assert len(firms) == len(vf.load_firms())
    assert firms["peak-xv"]["sources"] == 2            # Luma calendar + home page
    assert firms["inflexor"]["sources"] == 0 and firms["inflexor"]["no_events"]
    assert {s["key"] for s in upcoming["sources"]} >= {peak.key}


def test_api_filters_by_organiser_and_dates() -> None:
    now = utcnow()
    add_event(feed_for("peak-xv"), "Peak event", now + timedelta(days=3))
    add_event(feed_for("accel"), "Accel event", now + timedelta(days=5))

    client = TestClient(app)
    only = client.get("/api/insights/events", params={"organiser": "accel"}).json()
    assert [e["title"] for e in only["events"]] == ["Accel event"]
    assert only["counts"] == {"upcoming": 1, "past": 0}

    day = (now + timedelta(days=3)).astimezone(IST).date().isoformat()
    ranged = client.get("/api/insights/events", params={"start": day, "end": day}).json()
    assert [e["title"] for e in ranged["events"]] == ["Peak event"]

    assert client.get("/api/insights/events", params={"when": "someday"}).status_code == 400
    assert client.get("/api/insights/events", params={"organiser": "nobody"}).status_code == 404


def test_an_event_in_two_sources_shows_once_under_the_firm() -> None:
    starts = utcnow() + timedelta(days=4)
    add_event(feed_for("peak-xv"), "Shared", starts, url="https://luma.com/shared")
    pasted = ev.add_source("https://luma.com/some-community-calendar", label="Community")
    add_event(pasted, "Shared", starts, url="https://luma.com/shared")

    events = TestClient(app).get("/api/insights/events").json()["events"]
    assert [(e["title"], e["organiser_key"]) for e in events] == [("Shared", "peak-xv")]


def test_excel_has_the_same_rows() -> None:
    from openpyxl import load_workbook

    add_event(feed_for("peak-xv"), "Surge Open House", utcnow() + timedelta(days=2),
              city="Bengaluru", online=False)
    resp = TestClient(app).get("/api/insights/events.xlsx")
    assert resp.status_code == 200
    ws = load_workbook(io.BytesIO(resp.content)).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][:2] == ("Organiser", "Event")
    assert rows[1][0] == "Peak XV Partners" and rows[1][1] == "Surge Open House"
    assert rows[1][4] == "Bengaluru" and rows[1][6] == "In person"


def test_add_a_luma_link_reads_it_through_apify(monkeypatch) -> None:
    calls = stub_actor(monkeypatch, {"future": [LUMA_ROW], "past": []})
    client = TestClient(app)
    out = client.post("/api/insights/events/sources",
                      json={"url": "lu.ma/founders-club", "firm": "peak-xv"}).json()

    assert out["added"] and out["via"] == "luma" and out["error"] is None
    assert out["found"] == 1 and out["new"] == 1
    # A new calendar gets its upcoming events and its past ones.
    assert [c["input"]["period"] for c in calls] == ["future", "past"]
    sources = client.get("/api/insights/events").json()["sources"]
    pasted = next(s for s in sources if s["custom"])
    assert pasted["url"] == "https://lu.ma/founders-club" and pasted["firm"] == "peak-xv"
    assert pasted["label"] == "Peak XV Partners"


def test_add_a_website_reads_it_through_firecrawl(monkeypatch) -> None:
    day = (utcnow() + timedelta(days=20)).astimezone(IST).date().isoformat()
    stub_page(monkeypatch, f"Demo day on {day}", [{"title": "Demo Day", "start": day}])
    client = TestClient(app)
    out = client.post("/api/insights/events/sources",
                      json={"url": "https://angels.test/events",
                            "label": "Angels Network"}).json()
    assert out["added"] and out["via"] == "page"
    assert (out["new"], out["credits_used"]) == (1, 1)

    shown = client.get("/api/insights/events").json()["events"]
    assert [(e["title"], e["organiser"], e["via_label"], e["date_only"]) for e in shown] == [
        ("Demo Day", "Angels Network", "Website, via Firecrawl", True)]


@pytest.mark.parametrize("body, words", [
    ({"url": ""}, "Paste a link"),
    ({"url": "not a link"}, "does not look like"),
    ({"url": "ftp://x.test/events"}, "does not look like"),
    ({"url": "https://x.test/events", "firm": "nobody"}, "Unknown firm"),
    ({"url": "https://luma.com/peakxvpartners"}, "already a source"),
])
def test_add_refuses_what_it_cannot_use(monkeypatch, body, words) -> None:
    calls = stub_actor(monkeypatch, {})
    out = TestClient(app).post("/api/insights/events/sources", json=body).json()
    assert not out["added"] and words in out["error"]
    assert calls == []


def test_remove_a_pasted_source_with_its_events() -> None:
    feed = ev.add_source("https://angels.test/events", label="Angels")
    add_event(feed, "Pitch night", utcnow() + timedelta(days=3))
    client = TestClient(app)

    assert client.delete(f"/api/insights/events/sources/{feed.source_id}").json() == {"removed": True}
    assert client.get("/api/insights/events").json()["events"] == []
    assert client.delete(f"/api/insights/events/sources/{feed.source_id}").status_code == 404


# --- Refresh now ---------------------------------------------------------------------

def test_manual_refresh_has_an_events_scope(monkeypatch) -> None:
    monkeypatch.setattr(scheduler, "_last_event_read", lambda: None)
    monkeypatch.setattr(ev, "refresh_events",
                        lambda: ev.EventsResult(reads=[EventRead(items_new=2, items_seen=5)]))
    client = TestClient(app)

    started = client.post("/api/refresh", json={"scope": "events"}).json()
    assert started["started"] is True
    deadline = time.monotonic() + 5
    while scheduler.manual_status().running and time.monotonic() < deadline:
        time.sleep(0.02)
    status = client.get("/api/refresh").json()
    assert status["scope"] == "events"
    assert status["summary"] == "Refreshed events: 2 new events from 1 reads."
