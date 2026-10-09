"""LinkedIn tests. Offline: the Apify actor and the model are stubbed.

Protected: a pasted link is kept as one canonical profile or company page and
anything else is refused in words; nothing is saved while Apify is not set up;
the first read asks the actor for the newest posts only (no reactions or
comments, both caps set), later reads only for posts newer than the newest
stored; a post is stored once; rows that are not posts are never stored and
an empty first read is an error; the label becomes the account's own name;
the model reads each post once, Funding details are kept only on Funding, and
a post it skipped stays pending; an account is read at most every
`linkedin_hours` and never with --skip-paid; the API filters by account, kind
and dates, hides Other unless asked, exports the same rows to Excel, adds and
removes accounts; "Refresh now" has a linkedin scope.
"""

from __future__ import annotations

import io
import re
import time
from datetime import datetime, timedelta, timezone

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import delete, func, select

from app import scheduler
from app.config import settings
from app.db import session_scope
from app.ingest import linkedin as li
from app.llm.bedrock import LlmResult
from app.llm.budget import BudgetExceeded
from app.main import app
from app.models import LinkedinKind, LinkedinPost, LinkedinRead, LinkedinSource, utcnow
from app.search import apify

COMPANY = "https://www.linkedin.com/company/acme-ventures/"
MS = 1000


def row(post_id: str, content: str, when: datetime, author: str = "Acme Ventures",
        ident: str = "acme-ventures", **extra) -> dict:
    """One post as the actor documents it, trimmed."""
    return {
        "type": "post", "id": post_id, "content": content,
        "linkedinUrl": f"https://www.linkedin.com/posts/acme_{post_id}",
        "author": {"name": author, "universalName": ident, "publicIdentifier": None,
                   "type": "company"},
        "postedAt": {"timestamp": int(when.timestamp() * MS), "date": when.isoformat()},
        "engagement": {"likes": 3},
        **extra,
    }


NOW = datetime(2026, 10, 8, 6, 0, tzinfo=timezone.utc)
POSTS = [
    row("701", "Thrilled to lead Orbo's $2M seed round with Blume Ventures!", NOW),
    row("702", "We are hiring an associate. Apply below.", NOW - timedelta(days=2)),
    row("703", "Kite is launching its app in Pune today.", NOW - timedelta(days=40),
        author="Kite", ident="kite-app"),
]
ANSWERS = {
    "Thrilled": {"kind": "Funding", "headline": "Acme leads Orbo's $2M seed round",
                 "company": "Orbo", "round": "seed", "amount": "$2M",
                 "investors": ["Acme Ventures", "Blume Ventures"]},
    "We": {"kind": "Other", "headline": "Acme is hiring", "round": "Seed",
           "amount": "$9M", "investors": ["X"]},
    "Kite": {"kind": "News", "headline": "Kite launches in Pune", "company": "Kite"},
}


@pytest.fixture(autouse=True)
def _clean(monkeypatch):
    monkeypatch.setattr(li, "invoke_json", lambda *a, **k: pytest.fail("the model was called"))
    yield
    with session_scope() as s:
        s.execute(delete(LinkedinPost))
        s.execute(delete(LinkedinRead))
        s.execute(delete(LinkedinSource))


@pytest.fixture
def token(monkeypatch):
    monkeypatch.setattr(settings, "apify_api_token", "apify-test")


def stub_actor(monkeypatch, *answers: list[dict] | Exception) -> list[dict]:
    """The actor answers each run with the next list (or raises it). Returns
    every call made."""
    calls, queue = [], list(answers)

    def fake(actor, run_input, *, max_items, max_charge_usd=None):
        calls.append({"actor": actor, "input": run_input, "max_items": max_items,
                      "max_charge_usd": max_charge_usd})
        answer = queue.pop(0) if queue else []
        if isinstance(answer, Exception):
            raise answer
        return [dict(r) for r in answer]

    monkeypatch.setattr(li.apify, "run_actor", fake)
    return calls


def stub_model(monkeypatch, answers: dict[str, dict] = ANSWERS,
               fail: Exception | None = None) -> list[str]:
    """The model answers by the first word of each numbered post."""
    prompts = []

    def fake(feature, prompt, **kw):
        assert feature == li.FEATURE
        if fail:
            raise fail
        prompts.append(prompt)
        out, number = [], None
        for line in prompt.splitlines():
            m = re.match(r"\[(\d+)\] \(", line)
            if m:
                number = int(m.group(1))
                continue
            word = line.strip().split(" ", 1)[0] if number else None
            if number and word in answers:
                out.append({"i": number, **answers[word]})
            number = None
        return out, LlmResult("", settings.model_cheap, 900, 150, 0.0002)

    monkeypatch.setattr(li, "invoke_json", fake)
    return prompts


def add(url: str = COMPANY, label: str | None = None) -> LinkedinSource:
    return li.add_source(url, label)


# --- Links ------------------------------------------------------------------------

@pytest.mark.parametrize("pasted, url, kind", [
    ("https://www.linkedin.com/company/acme-ventures/", COMPANY, "company"),
    ("linkedin.com/company/Acme-Ventures/posts/?feedView=all", COMPANY, "company"),
    ("https://in.linkedin.com/in/jane-doe-12ab?trk=public", "https://www.linkedin.com/in/jane-doe-12ab/",
     "profile"),
    ("http://www.linkedin.com/in/Jane-Doe/recent-activity/all/",
     "https://www.linkedin.com/in/jane-doe/", "profile"),
])
def test_links_become_one_canonical_account(pasted, url, kind) -> None:
    got_url, got_kind, _ = li.canonical_url(pasted)
    assert (got_url, got_kind) == (url, kind)


@pytest.mark.parametrize("pasted, says", [
    ("", "Paste a LinkedIn link"),
    ("https://twitter.com/acme", "not a LinkedIn link"),
    ("https://evil.com/linkedin.com/in/x", "not a LinkedIn link"),
    ("https://linkedin.com.evil.com/in/x", "not a LinkedIn link"),
    ("javascript:alert(1)//linkedin.com/in/x", "not a LinkedIn link"),
    ("https://www.linkedin.com/posts/acme_activity-7000-abc", "single post"),
    ("https://www.linkedin.com/feed/update/urn:li:activity:7000/", "single post"),
    ("https://www.linkedin.com/school/iit-bombay/", "profile (linkedin.com/in/...)"),
])
def test_other_links_are_refused_in_words(pasted, says) -> None:
    with pytest.raises(li.SourceError, match=re.escape(says)):
        li.canonical_url(pasted)


def test_nothing_is_saved_while_apify_is_not_set_up(monkeypatch) -> None:
    monkeypatch.setattr(settings, "apify_api_token", "")
    with pytest.raises(li.SourceError, match="APIFY_API_TOKEN"):
        add()
    assert li.all_sources() == []


def test_the_same_account_cannot_be_added_twice(token) -> None:
    first = add()
    assert first.label == "acme-ventures" and first.label_auto and first.kind == "company"
    with pytest.raises(li.SourceError, match="already a source"):
        add("https://linkedin.com/company/ACME-VENTURES/about/")
    named = add("https://www.linkedin.com/in/jane-doe/", "  Jane   Doe ")
    assert named.label == "Jane Doe" and not named.label_auto


# --- Reading ----------------------------------------------------------------------

def test_first_read_takes_the_newest_posts_with_both_caps(token, monkeypatch) -> None:
    calls = stub_actor(monkeypatch, POSTS + [{"type": "notice", "message": "done"}])
    source = add()
    run = li.read_source(source)

    assert run.error is None and run.items_seen == 3 and run.items_new == 3
    call = calls[0]
    assert call["actor"] == settings.apify_linkedin_actor
    assert call["input"] == {
        "targetUrls": [COMPANY], "maxPosts": 20, "includeReposts": True,
        "includeQuotePosts": True, "scrapeReactions": False, "scrapeComments": False,
    }
    assert call["max_items"] == 20
    assert 0 < call["max_charge_usd"] <= settings.apify_max_charge_usd
    # Three posts charged, plus the run start; the notice row is free.
    assert run.apify_usd == pytest.approx(3 * 0.002 + 0.00005)

    with session_scope() as s:
        posts = {p.post_id: p for p in s.scalars(select(LinkedinPost)).all()}
        label = s.get(LinkedinSource, source.id).label
    assert set(posts) == {"701", "702", "703"}
    assert posts["701"].posted_at == NOW and not posts["701"].repost
    assert posts["703"].repost and posts["703"].author == "Kite"
    assert posts["701"].classified_at is None
    # The slug gave way to the account's own name, from its own post.
    assert label == "Acme Ventures"


def test_later_reads_ask_only_for_newer_posts_and_store_each_once(token, monkeypatch) -> None:
    newer = row("704", "Orbo crosses 1 million users.", NOW + timedelta(hours=5))
    calls = stub_actor(monkeypatch, POSTS, [POSTS[0], newer], [])
    source = add()
    li.read_source(source)
    second = li.read_source(source)
    quiet = li.read_source(source)

    assert calls[1]["input"]["postedLimitDate"] == "2026-10-08T06:00:01.000Z"
    assert second.items_seen == 2 and second.items_new == 1
    # Nothing new is not a failure; it is charged as an empty read.
    assert quiet.error is None and quiet.items_seen == 0
    assert quiet.apify_usd == pytest.approx(0.001 + 0.00005)
    with session_scope() as s:
        assert s.scalar(select(func.count(LinkedinPost.id))) == 4


def test_rows_that_are_not_posts_are_never_stored(token, monkeypatch) -> None:
    stub_actor(monkeypatch, [
        {"type": "post", "id": "1", "content": "No link at all"},
        {"type": "post", "id": "2", "content": "Bad link", "linkedinUrl": "javascript:alert(1)"},
        {"type": "post", "id": "3", "content": "   ", "linkedinUrl": "https://www.linkedin.com/posts/3"},
        {"type": "reaction", "id": "4", "linkedinUrl": "https://www.linkedin.com/posts/4"},
        row("5", "A real post about Orbo's seed round.", NOW),
    ])
    run = li.read_source(add())
    with session_scope() as s:
        assert list(s.scalars(select(LinkedinPost.post_id)).all()) == ["5"]
    assert run.items_new == 1


def test_an_empty_or_failed_first_read_is_an_error_not_a_crash(token, monkeypatch) -> None:
    stub_actor(monkeypatch, [],
               [{"error": "Profile not found", "query": COMPANY}],
               apify.ApifyError("Apify credits or plan limit used up. Check your Apify billing."))
    source = add()
    empty = li.read_source(source)
    problem = li.read_source(source)
    billing = li.read_source(source)
    assert "No posts found" in empty.error
    assert problem.error == "LinkedIn could not be read for that account: Profile not found"
    assert "Apify credits" in billing.error
    with session_scope() as s:
        assert s.scalar(select(func.count(LinkedinRead.id))) == 3


# --- The model --------------------------------------------------------------------

def test_the_model_reads_each_post_once(token, monkeypatch) -> None:
    stub_actor(monkeypatch, POSTS)
    source = add()
    li.read_source(source)
    prompts = stub_model(monkeypatch)

    result = li.classify_pending()
    assert result.classified == 3 and result.batches == 1
    assert "(Account: Acme Ventures)" in prompts[0]
    assert "(Account: Acme Ventures, reposting Kite)" in prompts[0]
    with session_scope() as s:
        posts = {p.post_id: p for p in s.scalars(select(LinkedinPost)).all()}
    funding, other, news = posts["701"], posts["702"], posts["703"]
    assert funding.kind == LinkedinKind.FUNDING and funding.round_label == "Seed"
    assert funding.investors == "Acme Ventures; Blume Ventures" and funding.amount == "$2M"
    # Round, amount and investors belong to Funding only.
    assert other.kind == LinkedinKind.OTHER and other.round_label is None
    assert other.amount is None and other.investors is None
    assert news.kind == LinkedinKind.NEWS and news.company == "Kite"

    assert li.classify_pending().considered == 0
    assert len(prompts) == 1


def test_a_skipped_post_stays_pending_and_the_budget_stops_cleanly(token, monkeypatch) -> None:
    stub_actor(monkeypatch, POSTS)
    li.read_source(add())
    stub_model(monkeypatch, {"Thrilled": ANSWERS["Thrilled"]})
    assert li.classify_pending().classified == 1

    stub_model(monkeypatch, fail=BudgetExceeded("daily cap reached", None))
    stopped = li.classify_pending()
    assert stopped.classified == 0 and "daily cap" in stopped.stopped_reason
    with session_scope() as s:
        assert s.scalar(select(func.count(LinkedinPost.id))
                        .where(LinkedinPost.classified_at.is_(None))) == 2


# --- One refresh --------------------------------------------------------------------

def test_an_account_is_read_at_most_every_linkedin_hours(token, monkeypatch) -> None:
    calls = stub_actor(monkeypatch, POSTS, [])
    stub_model(monkeypatch)
    source = add()
    first = li.refresh_linkedin()
    assert len(first.reads) == 1 and first.classify.classified == 3

    again = li.refresh_linkedin()
    assert again.reads == [] and again.skipped == 1 and len(calls) == 1

    with session_scope() as s:
        for r in s.scalars(select(LinkedinRead)).all():
            r.started_at = utcnow() - timedelta(hours=settings.linkedin_hours, minutes=1)
    assert li.is_due(source)
    assert len(li.refresh_linkedin().reads) == 1


def test_a_failed_read_does_not_count_and_skip_paid_reads_nothing(token, monkeypatch) -> None:
    calls = stub_actor(monkeypatch, apify.ApifyError("Apify took too long to answer."))
    source = add()
    li.refresh_linkedin(classify=False)
    assert li.is_due(source)

    paid = li.refresh_linkedin(free_only=True, classify=False)
    assert paid.reads == [] and paid.skipped == 1 and len(calls) == 1


# --- API ------------------------------------------------------------------------------

def seed(monkeypatch) -> LinkedinSource:
    stub_actor(monkeypatch, POSTS)
    stub_model(monkeypatch)
    source = add()
    li.read_now(source)
    return source


def test_api_lists_funding_and_news_with_filters(token, monkeypatch) -> None:
    source = seed(monkeypatch)
    client = TestClient(app)

    body = client.get("/api/linkedin").json()
    assert [p["kind"] for p in body["posts"]] == ["Funding", "News"]
    assert body["counts"] == {"funding": 1, "news": 1, "other": 1}
    assert body["hidden_other"] == 1 and body["pending"] == 0 and body["apify_ready"]
    top = body["posts"][0]
    assert top["account"] == "Acme Ventures" and top["investors"] == ["Acme Ventures", "Blume Ventures"]
    assert body["posts"][1]["repost"] and body["posts"][1]["author"] == "Kite"
    [src] = body["sources"]
    assert src["id"] == source.id and src["count"] == 2 and src["posts"] == 3
    assert src["kind_label"] == "Company page" and src["last_error"] is None

    assert len(client.get("/api/linkedin", params={"all": "true"}).json()["posts"]) == 3
    funding = client.get("/api/linkedin", params={"kind": "funding"}).json()
    assert [p["kind"] for p in funding["posts"]] == ["Funding"]
    # Dates are Indian days, inclusive: 8 Oct 2026 IST holds only the newest post.
    day = client.get("/api/linkedin", params={"start": "2026-10-08", "end": "2026-10-08"}).json()
    assert [p["headline"] for p in day["posts"]] == ["Acme leads Orbo's $2M seed round"]
    assert client.get("/api/linkedin", params={"kind": "jobs"}).status_code == 400
    assert client.get("/api/linkedin", params={"source": 999}).status_code == 404
    assert client.get("/api/linkedin", params={"start": "2026-10-09",
                                                "end": "2026-10-01"}).status_code == 400


def test_api_excel_has_the_same_rows(token, monkeypatch) -> None:
    from openpyxl import load_workbook

    seed(monkeypatch)
    resp = TestClient(app).get("/api/linkedin.xlsx", params={"kind": "funding"})
    assert resp.status_code == 200
    assert "gps-linkedin-funding" in resp.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(resp.content)).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][:3] == ("Account", "What", "Company")
    assert len(rows) == 2 and rows[1][1] == "Funding" and rows[1][2] == "Orbo"
    assert ws.cell(row=2, column=10).hyperlink.target.startswith("https://www.linkedin.com/posts/")


def test_api_adds_reads_and_removes_an_account(token, monkeypatch) -> None:
    stub_actor(monkeypatch, POSTS)
    stub_model(monkeypatch)
    client = TestClient(app)

    added = client.post("/api/linkedin/sources", json={"url": COMPANY}).json()
    assert added["added"] and added["label"] == "Acme Ventures"
    assert (added["found"], added["new"], added["funding"], added["news"]) == (3, 3, 1, 1)
    assert added["apify_usd"] > 0 and added["llm_usd"] == pytest.approx(0.0002)
    assert added["error"] is None

    refused = client.post("/api/linkedin/sources", json={"url": "https://x.com/acme"}).json()
    assert refused["added"] is False and "not a LinkedIn link" in refused["error"]

    assert client.delete(f"/api/linkedin/sources/{added['source_id']}").json() == {"removed": True}
    with session_scope() as s:
        assert s.scalar(select(func.count(LinkedinPost.id))) == 0
        assert s.scalar(select(func.count(LinkedinRead.id))) == 0
    assert client.delete(f"/api/linkedin/sources/{added['source_id']}").status_code == 404


def test_api_add_reports_a_model_that_was_stopped(token, monkeypatch) -> None:
    stub_actor(monkeypatch, POSTS)
    stub_model(monkeypatch, fail=BudgetExceeded("total cap reached", None))
    added = TestClient(app).post("/api/linkedin/sources", json={"url": COMPANY}).json()
    assert added["added"] and added["new"] == 3
    assert added["error"] == "The posts were not read yet: total cap reached"


def test_manual_refresh_has_a_linkedin_scope(monkeypatch) -> None:
    monkeypatch.setattr(scheduler, "_last_linkedin_read", lambda: None)
    monkeypatch.setattr(li, "refresh_linkedin", lambda: li.LinkedinResult(skipped=2))
    client = TestClient(app)
    started = client.post("/api/refresh", json={"scope": "linkedin"}).json()
    assert started["started"] is True
    deadline = time.monotonic() + 5
    while scheduler.manual_status().running and time.monotonic() < deadline:
        time.sleep(0.02)
    status = client.get("/api/refresh").json()
    assert status["scope"] == "linkedin"
    assert status["summary"].startswith("LinkedIn is up to date")
