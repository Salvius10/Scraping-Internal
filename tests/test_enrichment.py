"""Enrichment: leaders of funded startups and their contacts, through SignalHire.

Offline: SignalHire is stubbed (conftest refuses a real client).

Protected: a title holds a wanted designation only when it says so at *this*
company -- not an assistant to the CEO, not a co-founder of a side business,
not a same-named company; the best email (work) and phone (mobile) are kept;
the client sends the key, the synchronous lookup and both caps, and words
every failure, marking the ones that must stop the job; a company is sent
once, a person paid for once; a stop error leaves the queue waiting, a company
error does not stop the job, credits running out mid-company resumes cleanly;
the API refuses without a key, shows only what was found, and the Excel file
holds the page with a clickable LinkedIn link and no formula.
"""

from __future__ import annotations

import io
import time
from datetime import datetime, timezone

import httpx
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select

from app.api import funding as funding_api
from app.config import settings
from app.db import session_scope
from app.enrichment import roles, worker
from app.main import app
from app.models import EnrichedContact, EnrichmentCompany, EnrichStatus, FundedStartup, Stage
from app.proclock import enrichment_lock
from app.search import signalhire
from app.search.signalhire import Found, Lookup, Profile, SignalHireError

client = TestClient(app)


def uid(n: int) -> str:
    return f"{n:032d}"


def profile(n: int, name: str, *jobs: tuple[str, str],
            location: str = "Bengaluru, Karnataka, India") -> Profile:
    return Profile(uid=uid(n), full_name=name, location=location, jobs=list(jobs))


# --- Designations ----------------------------------------------------------------------

@pytest.mark.parametrize("title, held", [
    ("Co-Founder & CEO", ["Co-founder", "CEO"]),
    ("Founder & CTO", ["Co-founder", "CTO"]),
    ("Cofounder", ["Co-founder"]),
    ("Co-Founder | Product & Tech", ["Co-founder"]),
    ("Chief Technology Officer (CTO)", ["CTO"]),
    ("Chief Executive Officer", ["CEO"]),
    ("Chief Operating Officer", ["COO"]),
    ("VP of Engineering", ["VP Engineering"]),
    ("Vice President, Engineering", ["VP Engineering"]),
    ("SVP Engineering", ["VP Engineering"]),
    ("Head of Engineering", ["Head of Engineering"]),
    ("Head - Engineering", ["Head of Engineering"]),
    ("Engineering Head", ["Head of Engineering"]),
    ("Head of Product", ["Head of Product"]),
    ("Head of Products", ["Head of Product"]),
    ("Product Head", ["Head of Product"]),
])
def test_titles_that_hold_a_designation(title, held) -> None:
    assert roles.designations(title) == held


@pytest.mark.parametrize("title", [
    "Executive Assistant to Co-Founder & CEO and Project Manager, CEO's office",   # seen live
    "Founder's Office",
    "Chief of Staff to the CEO",
    "Advisor to the Founders",
    "Ex-CTO",
    "Former CEO",
    "Associate Director - Engineering",                                            # seen live
    "Head of Product Marketing",
    "Founding Engineer",
    "AVP Engineering",
    "Assistant Vice President Engineering",
    "Engineering Manager",
    "CEO Intern",
    "",
])
def test_titles_that_only_name_a_role(title) -> None:
    assert roles.designations(title) == []


@pytest.mark.parametrize("a, b, same", [
    ("Zepto", "Zepto", True),
    ("Zepto (KiranaKart Technologies)", "Zepto", True),
    ("Classplus Technologies Pvt. Ltd.", "Classplus", True),
    ("Kirana Pro", "KiranaPro", True),
    ("Sarvam AI", "Sarvam", True),
    ("Zepto | 10-minute delivery", "zepto", True),
    ("Zepto Logic Technologies Private Limited", "Zepto", False),                  # seen live
    ("Zepto Digital Labs", "Zepto", False),                                        # seen live
    ("Ola Electric", "Ola", False),
    ("", "Zepto", False),
])
def test_same_company(a, b, same) -> None:
    assert roles.same_company(a, b) is same


def test_search_name_is_the_core_of_the_name() -> None:
    assert roles.search_name("Classplus Technologies Pvt. Ltd.") == "classplus"
    assert roles.search_name("Tech Labs") == "tech"            # never reduced to nothing
    assert roles.search_name("!!!") == "!!!"


# Shaped like the live search for "Zepto" (2026-10-27).
ZEPTO = [
    profile(1, "Vinay Coo", ("Zepto", "Chief Operating Officer"), ("Zepto", "SVP Operations")),
    profile(2, "Aadit Ceo", ("Zepto", "Co-Founder & CEO"), ("KiranaKart", "Founder & CEO")),
    profile(3, "Kaivalya Cofounder", ("Zepto", "Co-Founder | Product & Tech"),
            ("KiranaKart", "Founder & CTO")),
    profile(4, "Associate Director", ("Zepto", "Associate Director - Engineering"),
            ("Voizo", "Co-Founder and CTO")),
    profile(5, "Other Zepto", ("Zepto Logic Technologies Private Limited", "Co-Founder")),
    profile(6, "Real Cto", ("Zepto", "Chief Technology Officer"), ("Various Startups", "Investor")),
    profile(7, "Side Founder", ("Yoga Chemistry", "Founder & CEO"), ("Zepto", "Director")),
    profile(8, "Project Associate", ("Zepto", "Project Associate"),
            ("Kazmi Construction Company", "Co-Founder")),
    profile(9, "Assistant", ("Zepto", "Executive Assistant to Co-Founder & CEO")),
    profile(10, "Digital Labs", ("Zepto Digital Labs", "Founder")),
    profile(2, "Aadit Ceo", ("Zepto", "Co-Founder & CEO")),                       # repeated hit
]


def test_leaders_are_people_holding_a_designation_at_this_company() -> None:
    leaders = roles.pick_leaders(ZEPTO, "Zepto", limit=6)
    assert [x.full_name for x in leaders] == ["Aadit Ceo", "Kaivalya Cofounder", "Real Cto",
                                               "Vinay Coo"]               # founders first
    assert leaders[0].designations == ["Co-founder", "CEO"]
    assert leaders[0].title == "Co-Founder & CEO"
    capped = roles.pick_leaders(ZEPTO, "Zepto", limit=2)
    assert [x.full_name for x in capped] == ["Aadit Ceo", "Kaivalya Cofounder"]
    assert roles.pick_leaders(ZEPTO, "Swiggy", limit=6) == []


CANDIDATE = {
    "uid": uid(2),
    "fullName": "Aadit Ceo",
    "locations": [{"name": "Mumbai, Maharashtra, India"}],
    "contacts": [
        {"type": "email", "value": "aadit@gmail.test", "rating": "100", "subType": "personal"},
        {"type": "email", "value": "aadit@zepto.test", "rating": 70, "subType": "work"},
        {"type": "email", "value": "a.ceo@zepto.test", "rating": 100, "subType": "work"},
        {"type": "phone", "value": "+91 80 1234 5678", "rating": 100, "subType": "work_phone"},
        {"type": "phone", "value": "+91  98765 43210", "rating": 70, "subType": "mobile"},
        {"type": "link", "value": "javascript:alert(1)"},
        {"type": "email", "value": "", "subType": "work"},
    ],
    "social": [{"type": "fb", "link": "https://facebook.com/aadit"},
               {"type": "li", "link": "https://www.linkedin.com/in/aadit"}],
    "experience": [{"position": "Co-Founder & CEO", "company": "Zepto", "current": True,
                    "companyUrl": "https://www.linkedin.com/company/zeptonow",
                    "website": "https://www.zeptonow.com"},
                   {"position": "Founder", "company": "KiranaKart", "current": False,
                    "companyUrl": "https://www.linkedin.com/company/kiranakart", "website": "-"}],
}


def test_details_keep_the_best_email_phone_and_linkedin() -> None:
    d = roles.details(CANDIDATE, "Zepto")
    assert d.email == "a.ceo@zepto.test"
    assert d.other_emails == ["aadit@zepto.test", "aadit@gmail.test"]
    assert d.phone == "+91 98765 43210"                 # mobile first, spacing tidied
    assert d.other_phones == ["+91 80 1234 5678"]
    assert d.linkedin == "https://www.linkedin.com/in/aadit"
    assert d.title == "Co-Founder & CEO" and d.location == "Mumbai, Maharashtra, India"
    assert d.company_page == "https://www.linkedin.com/company/zeptonow"   # this company's
    assert d.company_site == "https://www.zeptonow.com"
    assert roles.details(CANDIDATE, "KiranaKart").company_site is None     # "-" is no site


# Seen live for "Karta AI": SignalHire calls both of these "Karta".
@pytest.mark.parametrize("company, page, site, score", [
    ("Karta AI", "https://www.linkedin.com/company/getkarta", "http://www.getkarta.ai", 2),
    ("Karta AI", "https://www.linkedin.com/company/kartashop", None, 1),
    ("Zepto", "https://www.linkedin.com/company/zepto/", None, 3),
    ("Classplus Technologies", None, "https://classplusapp.com", 1),
    ("Zepto", "https://www.linkedin.com/company/swiggy", "https://swiggy.com", 0),
    ("Zepto", None, None, 0),
])
def test_page_score(company, page, site, score) -> None:
    assert roles.page_score(company, page, site) == score


def test_details_take_only_web_links_and_survive_odd_answers() -> None:
    odd = {"contacts": [{"type": "link", "value": "https://in.linkedin.com/in/x"}],
           "social": [{"type": "li", "link": "javascript:alert(1)"}],
           "experience": "n/a", "locations": None}
    d = roles.details(odd, "Zepto")
    assert d.linkedin == "https://in.linkedin.com/in/x"
    assert d.email is d.phone is d.title is d.location is None
    assert roles.details({}, "Zepto") == roles.Details()


# --- The client --------------------------------------------------------------------------

def mock(monkeypatch, handler) -> list[httpx.Request]:
    seen: list[httpx.Request] = []

    def record(request: httpx.Request) -> httpx.Response:
        seen.append(request)
        return handler(request)

    monkeypatch.setattr(settings, "signalhire_api_key", "test-key")
    monkeypatch.setattr(signalhire, "_client", lambda timeout=None: httpx.Client(
        transport=httpx.MockTransport(record)))
    return seen


def test_search_sends_the_queries_and_reads_the_profiles(monkeypatch) -> None:
    import json

    def handler(request):
        return httpx.Response(200, headers={"X-Credits-Left": "549"}, json={
            "requestId": 7, "total": 57, "scrollId": "abc",
            "profiles": [
                {"uid": uid(1), "fullName": " Aadit  Ceo ", "location": "Mumbai, India",
                 "experience": [{"company": "Zepto", "title": "Co-Founder & CEO"},
                                {"company": None, "title": None}], "skills": []},
                {"uid": "not-a-uid", "fullName": "Broken"},
                {"uid": uid(2), "fullName": ""},
                "junk",
            ]})

    seen = mock(monkeypatch, handler)
    found = signalhire.search_people(company='Ze"pto', title=roles.TITLE_QUERY,
                                     location="India", size=500)
    request = seen[0]
    assert request.method == "POST"
    assert request.url == "https://www.signalhire.com/api/v1/candidate/searchByQuery"
    assert request.headers["apikey"] == "test-key"
    body = json.loads(request.content)
    assert body == {"currentCompany": '"Ze pto"', "currentTitle": roles.TITLE_QUERY,
                    "location": "India", "size": 100}
    assert found.total == 57 and found.credits_left == 549
    assert found.profiles == [Profile(uid=uid(1), full_name="Aadit Ceo", location="Mumbai, India",
                                      jobs=[("Zepto", "Co-Founder & CEO")])]


def test_lookup_is_synchronous_and_needs_no_callback(monkeypatch) -> None:
    import json

    seen = mock(monkeypatch, lambda request: httpx.Response(200, headers={"X-Credits-Left": "548"},
        json=[{"item": uid(1), "status": "success", "candidate": CANDIDATE}, "junk"]))
    answer = signalhire.lookup([uid(1), uid(1), "bad"])
    body = json.loads(seen[0].content)
    assert seen[0].url.path == "/api/v1/candidate/search"
    assert body == {"items": [uid(1)], "withoutWaterfall": True}      # no callbackUrl
    assert answer.credits_left == 548 and answer.results[0]["status"] == "success"
    assert len(answer.results) == 1

    assert signalhire.lookup(["bad"]).results == [] and len(seen) == 1   # nothing sent
    with pytest.raises(SignalHireError, match="100"):
        signalhire.lookup([uid(n) for n in range(101)])


def test_credits(monkeypatch) -> None:
    seen = mock(monkeypatch, lambda r: httpx.Response(200, json={"credits": 549}))
    assert signalhire.credits() == 549 and seen[0].method == "GET"


@pytest.mark.parametrize("status, call, stop, words", [
    (401, "search", True, "SIGNALHIRE_API_KEY"),
    (402, "search", True, "daily search quota"),
    (402, "lookup", True, "credits are used up"),
    (403, "search", True, "disabled"),
    (406, "lookup", False, "refused the request"),
    (422, "search", False, "Unknown location"),
    (500, "search", False, "server error"),
])
def test_failures_are_worded_and_say_whether_to_stop(monkeypatch, status, call, stop,
                                                     words) -> None:
    mock(monkeypatch, lambda r: httpx.Response(status, json={"message": "Unknown location"}))
    with pytest.raises(SignalHireError, match=words) as caught:
        if call == "search":
            signalhire.search_people(company="Zepto", title="CEO", location="Mars", size=10)
        else:
            signalhire.lookup([uid(1)])
    assert caught.value.stop is stop


def test_rate_limits_are_waited_out(monkeypatch) -> None:
    answers = [httpx.Response(429, headers={"Retry-After": "2"}), httpx.Response(429),
               httpx.Response(200, json={"credits": 5})]
    waits: list[float] = []
    mock(monkeypatch, lambda r: answers.pop(0))
    monkeypatch.setattr(signalhire, "_sleep", waits.append)
    assert signalhire.credits() == 5
    assert waits == [2.0, settings.signalhire_retry_waits[1]]


def test_a_rate_limit_that_lasts_stops_the_job(monkeypatch) -> None:
    seen = mock(monkeypatch, lambda r: httpx.Response(429))
    with pytest.raises(SignalHireError, match="rate limit") as caught:
        signalhire.lookup([uid(1)])
    assert caught.value.stop and len(seen) == len(settings.signalhire_retry_waits) + 1


def test_network_failures(monkeypatch) -> None:
    def unreachable(request):
        raise httpx.ConnectError("no route")

    mock(monkeypatch, unreachable)
    with pytest.raises(SignalHireError, match="Could not reach") as caught:
        signalhire.credits()
    assert caught.value.stop

    def slow(request):
        raise httpx.ReadTimeout("slow")

    mock(monkeypatch, slow)
    with pytest.raises(SignalHireError, match="too long") as caught:
        signalhire.credits()
    assert not caught.value.stop

    mock(monkeypatch, lambda r: httpx.Response(200, text="<html>"))
    with pytest.raises(SignalHireError, match="not JSON"):
        signalhire.credits()


def test_no_key_no_request() -> None:
    with pytest.raises(SignalHireError, match="not set up") as caught:
        signalhire.lookup([uid(1)])           # conftest's client would fail the test
    assert caught.value.stop


# --- The queue and the job ---------------------------------------------------------------

def funded(company: str, stage: Stage = Stage.SEED, day: int = 10, **over) -> FundedStartup:
    values = dict(company=company, company_key=company.lower(), stage=stage, round_label=None,
                  amount="$1M", investors=[], sector="Fintech", hq_city="Pune",
                  summary=f"{company} raised.", headline=f"{company} raises",
                  source_url=f"https://news.test/{company.lower()}-{day}", source_label="News",
                  source_type="rss", sources=[],
                  published_at=datetime(2026, 10, day, 6, 0, tzinfo=timezone.utc))
    values.update(over)
    return FundedStartup(**values)


def candidate(name: str, company: str, title: str, *, email: str | None = None,
              phone: str | None = None, page: str | None = None,
              site: str | None = None) -> dict:
    contacts = []
    if email:
        contacts.append({"type": "email", "value": email, "rating": 100, "subType": "work"})
    if phone:
        contacts.append({"type": "phone", "value": phone, "rating": 100, "subType": "mobile"})
    handle = name.split()[0].lower()
    return {"fullName": name, "contacts": contacts,
            "social": [{"type": "li", "link": f"https://www.linkedin.com/in/{handle}"}],
            "experience": [{"position": title, "company": company, "current": True,
                            "companyUrl": page or f"https://www.linkedin.com/company/{company.lower()}",
                            "website": site or "-"}],
            "locations": [{"name": "Pune, India"}]}


class FakeSignalHire:
    """SignalHire as the job sees it: people per company, answers per uid."""

    def __init__(self, people: dict[str, list[Profile]], answers: dict[str, dict],
                 credits: int = 100) -> None:
        self.people, self.answers, self.credits = people, answers, credits
        self.searches: list[str] = []
        self.lookups: list[list[str]] = []
        self.search_error: dict[str, Exception] = {}

    def search_people(self, *, company, title, location, size):
        assert title == roles.TITLE_QUERY and location == "India"
        self.searches.append(company)
        if company in self.search_error:
            raise self.search_error[company]
        profiles = self.people.get(company, [])
        return Found(total=len(profiles), profiles=profiles, credits_left=self.credits)

    def lookup(self, uids):
        self.lookups.append(list(uids))
        before = self.credits
        results = []
        for item in uids:
            answer = self.answers.get(item, {"status": "failed"})
            if answer["status"] == "success" and answer["candidate"]["contacts"]:
                if self.credits <= 0:
                    answer = {"status": "credits_are_over"}
                else:
                    self.credits -= 1                # as seen live: no contacts, no charge
            results.append({"item": item, **answer})
        # As seen live, the header gives the balance from before this lookup's charge.
        return Lookup(results=results, credits_left=before)

    def balance(self, timeout=None):
        return self.credits


@pytest.fixture
def fake(monkeypatch, clean_db):
    monkeypatch.setattr(settings, "signalhire_api_key", "test-key")
    monkeypatch.setattr(settings, "signalhire_location", "India")
    sh = FakeSignalHire(
        people={
            "acme": [profile(11, "Asha Founder", ("Acme", "Founder & CEO")),
                     profile(12, "Ravi Cto", ("Acme Technologies", "CTO")),
                     profile(13, "Not Leader", ("Acme", "Engineering Manager"))],
            "beta": [profile(21, "Beta Coo", ("Beta", "Chief Operating Officer"))],
        },
        answers={
            uid(11): {"status": "success", "candidate": candidate(
                "Asha Founder", "Acme", "Founder & CEO", email="asha@acme.test",
                phone="+91 99999 00000")},
            uid(21): {"status": "success", "candidate": candidate(
                "Beta Coo", "Beta", "Chief Operating Officer", email="coo@beta.test")},
        })
    monkeypatch.setattr(signalhire, "search_people", sh.search_people)
    monkeypatch.setattr(signalhire, "lookup", sh.lookup)
    monkeypatch.setattr(signalhire, "credits", sh.balance)
    monkeypatch.setattr(worker, "_state", worker.JobState())
    # What the status endpoint shows: the last count a job heard, no extra call.
    monkeypatch.setattr(worker, "credits_left", lambda: worker.state().credits_left)
    worker._credits.clear()
    yield sh
    wait_until_done()
    worker._credits.clear()


def wait_until_done(timeout: float = 5.0) -> worker.JobState:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        state = worker.state()
        if not state.running:
            return state
        time.sleep(0.02)
    raise AssertionError("the enrichment job did not finish")


def add_funded(*rows: FundedStartup) -> None:
    with session_scope() as s:
        s.add_all(rows)


def send_all() -> worker.Plan:
    """What "Send data for enrichment" does with no filters."""
    with session_scope() as s:
        rows = s.scalars(funding_api.ordered(funding_api.query(funding_api.FundingFilter()))).all()
        return worker.enqueue(s, rows)


def run_job() -> worker.JobState:
    assert worker.start() is True
    return wait_until_done()


def company_rows() -> dict[str, EnrichmentCompany]:
    with session_scope() as s:
        return {c.company_key: c for c in s.scalars(select(EnrichmentCompany))}


def contact_rows() -> dict[str, EnrichedContact]:
    with session_scope() as s:
        return {c.full_name: c for c in s.scalars(select(EnrichedContact))}


def test_each_company_is_sent_once(fake) -> None:
    add_funded(funded("Acme", Stage.SERIES_A, day=12), funded("Acme", Stage.SEED, day=1),
               funded("Beta", day=5))
    sent = send_all()
    assert sent.companies == 2 and len(sent.new) == 2 and sent.max_credits == 12
    rows = company_rows()
    assert rows["acme"].stage == "Series A"                      # the newest round
    assert rows["acme"].status == EnrichStatus.QUEUED

    again = send_all()                                           # still waiting
    assert again.waiting == 2 and not again.new
    run_job()
    after = send_all()                                           # enriched: not again
    assert after.enriched == 2 and not after.new and not after.waiting


def test_the_job_finds_leaders_and_their_contacts(fake) -> None:
    add_funded(funded("Acme", day=12), funded("Beta", day=10))   # sent newest first
    send_all()
    state = run_job()

    assert fake.searches == ["acme", "beta"]
    assert fake.lookups == [[uid(11), uid(12)], [uid(21)]]       # the manager is not paid for
    acme = company_rows()["acme"]
    assert acme.status == EnrichStatus.DONE and acme.people_found == 2
    assert acme.credits_used == 1 and acme.note == "2 leaders, 1 with an email or phone."

    people = contact_rows()
    asha, ravi = people["Asha Founder"], people["Ravi Cto"]
    assert asha.designation == "Co-founder, CEO" and asha.rank == 0
    assert asha.email == "asha@acme.test" and asha.phone == "+91 99999 00000"
    assert asha.linkedin == "https://www.linkedin.com/in/asha"
    assert asha.lookup == "success" and asha.location == "Pune, India"
    assert ravi.designation == "CTO" and ravi.lookup == "failed" and ravi.email is None
    assert "Not Leader" not in people

    assert state.error is None and state.worked == 2 and state.people == 3
    assert state.message == ("Enrichment finished: 2 companies enriched, 3 leaders found, "
                             "2 SignalHire credits used (98 left).")


def test_people_found_without_contacts_cost_nothing(fake) -> None:
    fake.answers[uid(12)] = {"status": "success",
                             "candidate": candidate("Ravi Cto", "Acme", "CTO")}   # no contacts
    add_funded(funded("Acme"))
    send_all()
    state = run_job()
    acme = company_rows()["acme"]
    assert acme.status == EnrichStatus.DONE and acme.credits_used == 1      # Asha only
    ravi = contact_rows()["Ravi Cto"]
    assert ravi.lookup == "success" and ravi.email is ravi.phone is None
    assert ravi.linkedin == "https://www.linkedin.com/in/ravi"
    # Read afresh when the job ends: the lookup's header still said 100.
    assert state.credits_used == 1 and state.credits_left == 99


KARTA_AI = dict(page="https://www.linkedin.com/company/getkarta", site="http://www.getkarta.ai")
KARTA_SHOP = dict(page="https://www.linkedin.com/company/kartashop")


def karta(fake, pages: list[dict]) -> None:
    """Three Karta leaders, as seen live; each person's company page from `pages`."""
    people = [(51, "Vijay Ceo", "Co-Founder and CEO"), (52, "Abhi Cofounder", "Co-Founder"),
              (53, "Narayan Ceo", "Founder & CEO")]
    fake.people["karta"] = [profile(n, name, ("Karta", title)) for n, name, title in people]
    for (n, name, title), page in zip(people, pages):
        fake.answers[uid(n)] = {"status": "success", "candidate": candidate(
            name, "Karta", title, email=f"{name.split()[0].lower()}@karta.test", **page)}
    add_funded(funded("Karta AI"))
    send_all()
    run_job()


def test_people_at_a_same_named_company_are_marked(fake) -> None:
    karta(fake, [KARTA_AI, KARTA_AI, KARTA_SHOP])
    people = contact_rows()
    assert not people["Vijay Ceo"].maybe_other and not people["Abhi Cofounder"].maybe_other
    assert people["Narayan Ceo"].maybe_other                      # at kartashop, not getkarta.ai
    assert people["Vijay Ceo"].company_site == "http://www.getkarta.ai"
    assert people["Narayan Ceo"].company_page == "https://www.linkedin.com/company/kartashop"
    assert company_rows()["karta ai"].note.endswith(
        "1 person seems to be at another company with the same name, and marked so.")

    page = client.get("/api/enrichment").json()
    assert [r["full_name"] for r in page["rows"]] == ["Vijay Ceo", "Abhi Cofounder",
                                                      "Narayan Ceo"]          # marked last
    assert [r["maybe_other"] for r in page["rows"]] == [False, False, True]


def test_pages_that_match_equally_are_not_guessed_between(fake) -> None:
    karta(fake, [dict(page="https://www.linkedin.com/company/karta-one"),
                 dict(page="https://www.linkedin.com/company/karta-two"),
                 dict(page="https://www.linkedin.com/company/karta-two")])
    assert not any(c.maybe_other for c in contact_rows().values())
    assert company_rows()["karta ai"].note.endswith("check each person's company page.")


def test_a_company_with_no_leaders(fake) -> None:
    fake.people["gamma"] = [profile(31, "Someone", ("Gamma", "Engineer"))]
    add_funded(funded("Gamma"), funded("Delta"))
    send_all()
    run_job()
    rows = company_rows()
    assert rows["gamma"].status == EnrichStatus.NO_MATCH
    assert "None of the 1 profiles" in rows["gamma"].note
    assert rows["delta"].status == EnrichStatus.NO_MATCH
    assert rows["delta"].note == "SignalHire found no leaders of this company in India."
    assert fake.lookups == []                                    # nothing paid for


def test_a_person_is_never_paid_for_twice(fake) -> None:
    fake.answers[uid(12)] = {"status": "timeout_exceeded"}
    add_funded(funded("Acme"))
    send_all()
    run_job()
    acme = company_rows()["acme"]
    assert acme.status == EnrichStatus.FAILED and "1 of 2 people" in acme.note

    fake.answers[uid(12)] = {"status": "success", "candidate": candidate(
        "Ravi Cto", "Acme", "CTO", email="ravi@acme.test")}
    with session_scope() as s:
        assert worker.requeue_failed(s) == 1
    run_job()
    assert fake.lookups[-1] == [uid(12)]                         # Asha is reused, not re-bought
    assert company_rows()["acme"].status == EnrichStatus.DONE
    people = contact_rows()
    assert people["Ravi Cto"].email == "ravi@acme.test"
    assert people["Asha Founder"].email == "asha@acme.test"     # kept


def test_a_stop_error_leaves_the_queue_waiting(fake) -> None:
    fake.search_error["acme"] = SignalHireError("SignalHire's daily search quota is used up.",
                                                stop=True)
    add_funded(funded("Acme", day=2), funded("Beta", day=3), funded("Gamma", day=1))
    send_all()                                                   # Beta, then Acme, then Gamma
    state = run_job()
    assert fake.searches == ["beta", "acme"]                     # stopped: Gamma not tried
    rows = company_rows()
    assert rows["beta"].status == EnrichStatus.DONE
    assert rows["acme"].status == rows["gamma"].status == EnrichStatus.QUEUED
    assert rows["acme"].note.startswith("Waiting: SignalHire's daily search quota")
    assert state.error.startswith("SignalHire's daily search quota")
    assert state.message.startswith("Enrichment stopped:")

    del fake.search_error["acme"]                                # the quota resets
    run_job()
    assert fake.searches[2:] == ["acme", "gamma"]
    assert company_rows()["acme"].status == EnrichStatus.DONE


def test_a_company_error_does_not_stop_the_job(fake) -> None:
    fake.search_error["acme"] = SignalHireError("SignalHire had a server error (HTTP 500).")
    add_funded(funded("Acme"), funded("Beta"))
    send_all()
    state = run_job()
    rows = company_rows()
    assert rows["acme"].status == EnrichStatus.FAILED and "HTTP 500" in rows["acme"].note
    assert rows["beta"].status == EnrichStatus.DONE
    assert state.error is None and state.worked == 2


def test_a_fault_fails_the_company_and_stops(fake) -> None:
    fake.search_error["acme"] = ValueError("bug")
    fake.search_error["beta"] = ValueError("bug")
    add_funded(funded("Acme"), funded("Beta"))
    send_all()
    state = run_job()
    assert len(fake.searches) == 1                              # not repeated for every company
    assert company_rows()[fake.searches[0]].status == EnrichStatus.FAILED
    assert "ValueError" in state.error


def test_credits_running_out_mid_company(fake) -> None:
    fake.credits = 1
    fake.answers[uid(12)] = {"status": "success", "candidate": candidate("Ravi Cto", "Acme", "CTO",
                                                                         email="ravi@acme.test")}
    add_funded(funded("Acme"))
    send_all()
    state = run_job()
    acme = company_rows()["acme"]
    assert acme.status == EnrichStatus.QUEUED and "credits ran out" in acme.note
    assert "credits ran out" in state.error
    assert contact_rows()["Asha Founder"].email == "asha@acme.test"     # kept

    fake.credits = 10                                            # topped up
    run_job()
    assert fake.lookups[-1] == [uid(12)]
    assert company_rows()["acme"].status == EnrichStatus.DONE
    assert contact_rows()["Ravi Cto"].email == "ravi@acme.test"


def test_an_unreadable_balance_never_keeps_the_job_open(fake, monkeypatch) -> None:
    def broken(timeout=None):
        raise RuntimeError("odd answer")

    monkeypatch.setattr(signalhire, "credits", broken)
    add_funded(funded("Acme"))
    send_all()
    state = run_job()                                   # finishes, lock released
    assert state.error is None and company_rows()["acme"].status == EnrichStatus.DONE
    assert state.credits_left == 100                    # the last count heard
    assert run_job().worked == 0                         # a new job can start


def test_abandoned_work_is_picked_up_again(fake) -> None:
    add_funded(funded("Acme"))
    send_all()
    with session_scope() as s:
        s.query(EnrichmentCompany).update({"status": EnrichStatus.WORKING})
    run_job()
    assert company_rows()["acme"].status == EnrichStatus.DONE


def test_one_job_at_a_time(fake) -> None:
    other = enrichment_lock()
    assert other.acquire()
    try:
        assert worker.start() is False                          # another process is enriching
    finally:
        other.release()
    add_funded(funded("Acme"))
    send_all()
    assert run_job().worked == 1


# --- The API -------------------------------------------------------------------------

def test_send_refuses_without_a_key(clean_db) -> None:
    add_funded(funded("Acme"))
    preview = client.get("/api/enrichment/preview").json()
    assert preview["ready"] is False and preview["credits_left"] is None and preview["new"] == 1
    answer = client.post("/api/enrichment").json()
    assert answer["sent"] is False and "SIGNALHIRE_API_KEY" in answer["message"]
    assert company_rows() == {}
    assert client.post("/api/enrichment/resume").json()["started"] is False


def test_preview_says_what_it_would_cost(fake, monkeypatch) -> None:
    monkeypatch.setattr(worker, "credits_left", lambda: 549)
    add_funded(funded("Acme"), funded("Beta", stage=Stage.SERIES_A))
    body = client.get("/api/enrichment/preview").json()
    assert body == {"ready": True, "companies": 2, "new": 2, "waiting": 0, "enriched": 0,
                    "over_cap": 0, "max_credits": 12, "per_company": 6, "credits_left": 549}
    seed = client.get("/api/enrichment/preview", params={"stage": "seed"}).json()
    assert seed["companies"] == 1
    assert client.get("/api/enrichment/preview", params={"stage": "z"}).status_code == 400
    assert company_rows() == {}                                  # a preview writes nothing


def test_send_caps_the_companies_per_click(fake, monkeypatch) -> None:
    monkeypatch.setattr(settings, "signalhire_max_companies", 1)
    add_funded(funded("Acme", day=2), funded("Beta", day=1))
    answer = client.post("/api/enrichment").json()
    assert answer["sent"] and answer["queued"] == 1 and answer["over_cap"] == 1
    assert "at most 1 are sent at a time" in answer["message"]
    wait_until_done()
    assert set(company_rows()) == {"acme"}                       # the newest


def test_send_runs_and_the_page_shows_what_was_found(fake, monkeypatch) -> None:
    monkeypatch.setattr(worker, "credits_left", lambda: 98)
    add_funded(funded("Acme"), funded("Beta"),
               funded('=HYPERLINK("http://evil")', stage=Stage.OTHER))
    assert client.post("/api/enrichment", params={"q": "zzz"}).json()["sent"] is False

    answer = client.post("/api/enrichment", params={"stage": "seed"}).json()
    assert answer["sent"] is True and answer["queued"] == 2
    assert answer["message"].startswith("Sent 2 companies for enrichment")
    wait_until_done()

    page = client.get("/api/enrichment").json()
    assert page["total"] == 3
    assert [r["full_name"] for r in page["rows"] if r["company"] == "Acme"] == [
        "Asha Founder", "Ravi Cto"]                              # founders first
    asha = next(r for r in page["rows"] if r["full_name"] == "Asha Founder")
    assert asha["email"] == "asha@acme.test" and asha["stage"] == "Seed"
    counts = {d["key"]: d["count"] for d in page["designations"]}
    assert counts == {"cofounder": 1, "ceo": 1, "cto": 1, "coo": 1, "vp-engineering": 0,
                      "head-engineering": 0, "head-product": 0}
    assert page["all_designations"] == 3            # not 4: Asha is Co-founder and CEO
    cto = client.get("/api/enrichment", params={"designation": "cto"}).json()
    assert [r["full_name"] for r in cto["rows"]] == ["Ravi Cto"]
    assert cto["total"] == 1 and cto["all_designations"] == 3
    assert {d["key"]: d["count"] for d in cto["designations"]}["coo"] == 1   # counts ignore the tab
    assert client.get("/api/enrichment", params={"q": "beta.test"}).json()["total"] == 1
    assert client.get("/api/enrichment", params={"designation": "cfo"}).status_code == 400

    status = client.get("/api/enrichment/status").json()
    assert status["ready"] and not status["running"]
    assert status["companies"] == 2 and status["done"] == 2 and status["waiting"] == 0
    assert status["contacts"] == 3 and status["reachable"] == 2
    assert status["credits_used"] == 2 and status["credits_left"] == 98
    assert status["message"].startswith("Enrichment finished")
    assert client.get("/api/status").json()["enriched_contacts"] == 3

    listed = client.get("/api/enrichment/companies").json()
    assert {c["company"]: c["status"] for c in listed} == {"Acme": "done", "Beta": "done"}

    again = client.post("/api/enrichment", params={"stage": "seed"}).json()
    assert again["sent"] and again["queued"] == 0 and again["enriched"] == 2
    assert "enriched before and not sent again" in again["message"]


def test_clear_takes_everything_off_the_list_but_keeps_what_was_paid_for(fake) -> None:
    from openpyxl import load_workbook

    add_funded(funded("Acme", day=12), funded("Beta", day=10))
    client.post("/api/enrichment")
    wait_until_done()
    assert client.get("/api/enrichment").json()["total"] == 3
    paid = fake.credits

    answer = client.delete("/api/enrichment").json()
    assert answer["cleared"] is True and answer["companies"] == 2
    assert answer["message"].startswith("Cleared 2 companies from the list.")
    assert answer["status"]["companies"] == answer["status"]["contacts"] == 0
    assert answer["status"]["message"] is None                  # the last job's outcome, gone
    assert client.get("/api/enrichment").json()["total"] == 0
    assert client.get("/api/enrichment/companies").json() == []
    assert client.get("/api/status").json()["enriched_contacts"] == 0
    ws = load_workbook(io.BytesIO(client.get("/api/export/enrichment.xlsx").content)).active
    assert ws.max_row == 1                                       # the header alone
    assert client.get("/api/enrichment/preview").json()["new"] == 2   # sendable again

    fake.lookups.clear()
    again = client.post("/api/enrichment").json()
    assert again["sent"] and again["queued"] == 2
    wait_until_done()
    assert fake.lookups == [[uid(12)]]          # only the one never found; Asha and Beta reused
    assert fake.credits == paid                 # nothing paid twice
    assert client.get("/api/enrichment").json()["total"] == 3
    assert {c["company"]: c["credits_used"] for c in
            client.get("/api/enrichment/companies").json()} == {"Acme": 0, "Beta": 0}


def test_clear_waits_for_a_running_job(fake) -> None:
    add_funded(funded("Acme"))
    send_all()
    worker._state.running = True                  # as if a job were mid-company
    try:
        answer = client.delete("/api/enrichment").json()
    finally:
        worker._state.running = False
    assert answer["cleared"] is False and "running" in answer["message"]
    assert company_rows()["acme"].cleared_at is None


def test_cleared_companies_leave_the_queue(fake) -> None:
    fake.search_error["beta"] = SignalHireError("SignalHire had a server error (HTTP 500).")
    add_funded(funded("Beta"))
    send_all()
    run_job()                                                    # Beta failed
    add_funded(funded("Acme"))
    send_all()                                                   # Acme waiting
    assert worker.clear() == 2
    assert run_job().worked == 0 and fake.searches == ["beta"]   # Acme not worked
    with session_scope() as s:
        assert worker.requeue_failed(s) == 0                      # Beta not retried


def test_init_db_adds_a_column_an_older_table_lacks(clean_db) -> None:
    import sqlite3

    from app.db import engine, init_db

    if sqlite3.sqlite_version_info < (3, 35):
        pytest.skip("SQLite too old to drop a column for the test")
    with engine.begin() as conn:
        conn.exec_driver_sql("ALTER TABLE enrichment_companies DROP COLUMN cleared_at")
    init_db()
    init_db()                                                    # idempotent
    with engine.begin() as conn:
        columns = {row[1] for row in conn.exec_driver_sql("PRAGMA table_info(enrichment_companies)")}
    assert "cleared_at" in columns


def test_retry_and_resume(fake) -> None:
    fake.search_error["acme"] = SignalHireError("SignalHire had a server error (HTTP 500).")
    add_funded(funded("Acme"))
    client.post("/api/enrichment")
    wait_until_done()
    assert company_rows()["acme"].status == EnrichStatus.FAILED
    resumed = client.post("/api/enrichment/resume").json()
    assert resumed["started"] is False and resumed["message"] == "Nothing is waiting to be enriched."
    assert resumed["status"]["failed"] == 1

    del fake.search_error["acme"]
    answer = client.post("/api/enrichment/retry").json()
    assert answer["started"] is True and answer["message"].startswith("Retrying 1 company.")
    wait_until_done()
    assert company_rows()["acme"].status == EnrichStatus.DONE


def test_enrichment_excel_holds_the_page_and_runs_no_formula(fake) -> None:
    from openpyxl import load_workbook

    # Searched as its core name, "hyperlink": the bracket is an aside.
    fake.people["hyperlink"] = [profile(41, '=HYPERLINK("http://evil","x")',
                                        ('=HYPERLINK("http://evil")', "CEO"))]
    add_funded(funded("Acme"), funded('=HYPERLINK("http://evil")'))
    client.post("/api/enrichment")
    wait_until_done()

    response = client.get("/api/export/enrichment.xlsx")
    assert response.status_code == 200
    assert "enriched-contacts-all-designations" in response.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(response.content)).active
    header = [c.value for c in ws[1]]
    assert header[:6] == ["Name", "Designation", "Company", "Email", "Phone", "LinkedIn"]
    rows = list(ws.iter_rows(min_row=2))
    assert len(rows) == 3
    asha = next(r for r in rows if r[0].value == "Asha Founder")
    assert asha[3].value == "asha@acme.test" and asha[4].value == "+91 99999 00000"
    assert asha[5].hyperlink.target == asha[5].value == "https://www.linkedin.com/in/asha"
    assert asha[header.index("Company LinkedIn page")].value == "https://www.linkedin.com/company/acme"
    assert asha[header.index("Check")].value is None                  # not marked
    assert isinstance(asha[header.index("Enriched (IST)")].value, datetime)
    evil = next(r for r in rows if str(r[0].value).startswith("=HYPERLINK"))
    assert evil[0].data_type == "s" and evil[2].data_type == "s"      # text, not formulas

    cto = load_workbook(io.BytesIO(client.get("/api/export/enrichment.xlsx",
                                              params={"designation": "cto"}).content)).active
    assert [r[0].value for r in cto.iter_rows(min_row=2)] == ["Ravi Cto", None]  # + no-CTO company


def test_enrichment_excel_lists_every_company_sent(fake) -> None:
    from openpyxl import load_workbook

    fake.search_error["delta"] = SignalHireError("SignalHire had a server error (HTTP 500).")
    add_funded(funded("Acme"), funded("Beta"), funded("Gamma"), funded("Delta"))
    client.post("/api/enrichment")
    wait_until_done()

    def export(**params) -> tuple[list[str], list[tuple]]:
        ws = load_workbook(io.BytesIO(client.get("/api/export/enrichment.xlsx",
                                                 params=params).content)).active
        header = [c.value for c in ws[1]]
        return header, list(ws.iter_rows(min_row=2))

    header, rows = export()
    name, company, status, check, stage = (header.index(h) for h in (
        "Name", "Company", "Status", "Check", "Funding stage"))
    seen = [(r[name].value, r[company].value, r[status].value) for r in rows]
    assert set(seen[:3]) == {("Asha Founder", "Acme", "Enriched"), ("Ravi Cto", "Acme", "Enriched"),
                             ("Beta Coo", "Beta", "Enriched")}
    assert seen[3:] == [(None, "Gamma", "No leaders found"), (None, "Delta", "Failed")]  # below
    gamma, delta = rows[3], rows[4]
    assert gamma[check].value == "SignalHire found no leaders of this company in India."
    assert "HTTP 500" in delta[check].value and gamma[stage].value == "Seed"
    assert gamma[company].font.italic and not rows[0][company].font.italic       # set apart

    # A designation tab: every company without one is listed too, and says so.
    header, rows = export(designation="cto")
    assert [(r[name].value, r[company].value, r[status].value) for r in rows] == [
        ("Ravi Cto", "Acme", "Enriched"), (None, "Gamma", "No leaders found"),
        (None, "Delta", "Failed"), (None, "Beta", "No CTO found")]
    assert rows[-1][check].value == "Enriched, but SignalHire found no CTO here."

    # A search narrows the companies by name, as it does the leaders.
    header, rows = export(q="gam")
    assert [(r[name].value, r[company].value) for r in rows] == [(None, "Gamma")]
