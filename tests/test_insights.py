"""Insights tests: funding-round extraction and the rounds API. Offline.

The model's output is untrusted: stages are validated, placeholders dropped,
and a story the model skipped is retried rather than stored empty. Each tab
shows only canonical Funding stories, newest first, and the Excel file holds
the same rows as the screen.
"""

from __future__ import annotations

import io
from datetime import timedelta

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import delete, select

from app.config import settings
from app.db import session_scope
from app.ingest import rounds
from app.llm.bedrock import LlmResult
from app.llm.budget import BudgetExceeded, get_spend
from app.main import app
from app.models import Article, Category, FundingRound, Stage, utcnow

SOURCE = "test_insights"


@pytest.fixture
def funding():
    """Three canonical Funding stories, one duplicate, one non-Funding story."""
    now = utcnow()
    specs = [
        ("ONYA raises Rs 12.5 crore in pre-Series A led by Divis", Category.FUNDING, 2, True),
        ("Rio Health raises $4.5 Mn Series A from Accel and Elevation", Category.FUNDING, 1, True),
        ("MyRx raises seed round led by SteerX VC", Category.FUNDING, 5, True),
        ("Rio Health bags $4.5 million", Category.FUNDING, 1, False),        # duplicate
        ("Zepto launches a new cafe", Category.PRODUCT_LAUNCH, 1, True),     # not funding
    ]
    ids = []
    with session_scope() as s:
        first_rio = None
        for i, (headline, category, days_ago, canonical) in enumerate(specs):
            article = Article(
                url=f"https://i.test/{i}", url_hash=f"i{i}", source=SOURCE,
                headline=headline, category=category,
                published_at=now - timedelta(days=days_ago),
                canonical_id=None if canonical else first_rio,
            )
            s.add(article)
            s.flush()
            if i == 1:
                first_rio = article.id
            ids.append(article.id)
    yield ids
    with session_scope() as s:
        s.execute(delete(FundingRound).where(FundingRound.article_id.in_(ids)))
        s.execute(delete(Article).where(Article.id.in_(ids)))


def stub_model(monkeypatch, answer):
    calls = []

    def fake(feature, prompt, **kw):
        calls.append(prompt)
        if isinstance(answer, Exception):
            raise answer
        return answer(prompt) if callable(answer) else answer, LlmResult(
            "", settings.model_cheap, 900, 300, 0.0003)

    monkeypatch.setattr(rounds, "invoke_json", fake)
    return calls


def _numbers(prompt: str) -> dict[str, int]:
    """Map each story's first word to its number in the prompt."""
    out = {}
    for line in prompt.splitlines():
        if line.startswith("[") and "]" in line:
            n, text = line[1:].split("]", 1)
            out[text.strip().split()[0]] = int(n)
    return out


ANSWERS = {
    "ONYA": {"company": "ONYA", "round": "Pre-Series A", "stage": "Seed",
             "amount": "Rs 12.5 crore", "investors": ["Divis"]},
    "Rio": {"company": "Rio Health", "round": "Series A", "stage": "Series A",
            "amount": "$4.5 Mn", "investors": ["Accel", "Elevation", "Accel"]},
    "MyRx": {"company": "MyRx", "round": "Seed", "stage": "Seed",
             "amount": "undisclosed", "investors": "SteerX VC"},
}


def answer_all(prompt):
    return [{"i": n, **ANSWERS[word]} for word, n in _numbers(prompt).items()
            if word in ANSWERS]


# --- Coercion -------------------------------------------------------------------

@pytest.mark.parametrize("stage, label, expected", [
    ("Seed", None, Stage.SEED),
    ("series a", None, Stage.SERIES_A),
    ("pre-seed", None, Stage.PRE_SEED),
    ("Series C", None, Stage.OTHER),
    ("nonsense", "Pre-Series A", Stage.SEED),      # falls back to the round name
    (None, "Series B extension", Stage.SERIES_B),
    (None, "Angel round", Stage.PRE_SEED),
    (None, "Debt", Stage.OTHER),
])
def test_stages_are_validated(stage, label, expected) -> None:
    assert rounds.coerce_stage(stage, label) == expected


def test_investors_and_placeholders_are_cleaned() -> None:
    assert rounds.clean_investors(["Accel", " accel ", "N/A", "Elevation"]) == "Accel; Elevation"
    assert rounds.clean_investors("Accel, Elevation and Nexus") == "Accel; Elevation; Nexus"
    assert rounds.clean_investors([]) is None
    assert rounds.clean_text("Undisclosed", 50) is None


# --- Extraction --------------------------------------------------------------------

def test_only_canonical_funding_stories_are_read(monkeypatch, funding) -> None:
    calls = stub_model(monkeypatch, answer_all)
    result = rounds.extract_pending()
    assert result.extracted >= 3
    prompt = "\n".join(calls)
    assert "Zepto launches" not in prompt           # not a Funding story
    assert "Rio Health bags" not in prompt          # duplicate


def test_nothing_is_paid_for_twice(monkeypatch, funding) -> None:
    calls = stub_model(monkeypatch, answer_all)
    rounds.extract_pending()
    first = len(calls)
    rounds.extract_pending()
    assert len(calls) == first


def test_a_skipped_story_is_retried_not_stored_empty(monkeypatch, funding) -> None:
    stub_model(monkeypatch, lambda p: [e for e in answer_all(p)
                                       if e["company"] != "MyRx"])
    rounds.extract_pending()
    with session_scope() as s:
        stored = {fr.company for fr in s.scalars(select(FundingRound)).all()}
    assert "MyRx" not in stored
    pending = [p.headline for p in rounds._load_pending(None)]
    assert any("MyRx" in h for h in pending)


def test_budget_exhaustion_stops_cleanly(monkeypatch, funding) -> None:
    stub_model(monkeypatch, BudgetExceeded("daily cap reached", get_spend()))
    result = rounds.extract_pending()
    assert result.extracted == 0 and "daily cap" in result.stopped_reason


# --- API ---------------------------------------------------------------------------

def test_rounds_by_stage_newest_first(monkeypatch, funding) -> None:
    stub_model(monkeypatch, answer_all)
    rounds.extract_pending()
    with TestClient(app) as client:
        seed = client.get("/api/insights/rounds", params={"stage": "seed"}).json()
        series_a = client.get("/api/insights/rounds", params={"stage": "series-a"}).json()

    ours = [r for r in seed["rounds"] if r["source"] == SOURCE]
    assert [r["company"] for r in ours] == ["ONYA", "MyRx"]    # 2 days ago, then 5
    onya = ours[0]
    assert onya["round"] == "Pre-Series A" and onya["amount"] == "Rs 12.5 crore"
    assert onya["investors"] == ["Divis"] and onya["published_at"]
    assert ours[1]["amount"] is None                           # "undisclosed" dropped

    rio = [r for r in series_a["rounds"] if r["source"] == SOURCE]
    assert rio[0]["investors"] == ["Accel", "Elevation"]
    counts = {c["key"]: c["count"] for c in seed["stages"]}
    assert list(counts) == ["pre-seed", "seed", "series-a", "series-b", "other"]
    assert counts["seed"] >= 2 and counts["series-a"] >= 1


def test_unknown_stage_is_a_404() -> None:
    with TestClient(app) as client:
        assert client.get("/api/insights/rounds", params={"stage": "series-z"}).status_code == 404


def test_excel_matches_the_screen(monkeypatch, funding) -> None:
    from openpyxl import load_workbook

    stub_model(monkeypatch, answer_all)
    rounds.extract_pending()
    with TestClient(app) as client:
        resp = client.get("/api/insights/rounds.xlsx", params={"stage": "seed"})
        screen = client.get("/api/insights/rounds", params={"stage": "seed"}).json()

    assert resp.status_code == 200
    assert "gps-funding-seed-" in resp.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(resp.content))["Seed"]
    header = [c.value for c in ws[1]]
    assert header[:6] == ["Company", "Stage", "Round", "Investors", "Amount", "Published (IST)"]
    assert ws.max_row - 1 == len(screen["rounds"])
    companies = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
    assert companies == [r["company"] or "" for r in screen["rounds"]]


def test_excel_without_a_stage_has_a_sheet_per_stage() -> None:
    from openpyxl import load_workbook

    with TestClient(app) as client:
        resp = client.get("/api/insights/rounds.xlsx")
    names = load_workbook(io.BytesIO(resp.content)).sheetnames
    assert names == ["Pre-Seed", "Seed", "Series A", "Series B", "Other rounds"]


@pytest.mark.parametrize("raw, tidy", [
    ("pre seed", "Pre-Seed"),
    ("pre-seed", "Pre-Seed"),
    ("pre-Series A", "Pre-Series A"),
    ("series a1", "Series A1"),
    ("bridge", "Bridge"),
    ("Pre-IPO", "Pre-IPO"),
    ("Series B", "Series B"),
    (None, None),
])
def test_round_names_have_one_spelling(raw, tidy) -> None:
    assert rounds.tidy_round(raw) == tidy


# --- Date range ------------------------------------------------------------------------

def _ours(body):
    return [r["company"] for r in body["rounds"] if r["source"] == SOURCE]


def test_date_range_filters_rows_and_counts(monkeypatch, funding) -> None:
    stub_model(monkeypatch, answer_all)
    rounds.extract_pending()
    from app.api.insights import IST
    today = utcnow().astimezone(IST).date()
    with TestClient(app) as client:
        everything = client.get("/api/insights/rounds", params={"stage": "seed"}).json()
        recent = client.get("/api/insights/rounds", params={
            "stage": "seed", "start": str(today - timedelta(days=3))}).json()
        older = client.get("/api/insights/rounds", params={
            "stage": "seed", "end": str(today - timedelta(days=4))}).json()

    assert _ours(everything) == ["ONYA", "MyRx"]
    assert _ours(recent) == ["ONYA"]                 # 2 days ago is in, 5 is out
    assert "MyRx" in _ours(older) and "ONYA" not in _ours(older)
    seed_count = lambda body: next(c["count"] for c in body["stages"] if c["key"] == "seed")
    assert seed_count(recent) < seed_count(everything)


def test_range_days_are_indian_days() -> None:
    from datetime import date
    from app.api.insights import Window

    w = Window(date(2026, 9, 25), date(2026, 9, 25))
    # 25 Sept IST runs from 24 Sept 18:30 UTC to 25 Sept 18:30 UTC.
    assert w.since.isoformat() == "2026-09-24T18:30:00+00:00"
    assert w.until.isoformat() == "2026-09-25T18:30:00+00:00"


def test_start_after_end_is_rejected() -> None:
    with TestClient(app) as client:
        resp = client.get("/api/insights/rounds", params={
            "stage": "seed", "start": "2026-09-25", "end": "2026-09-01"})
    assert resp.status_code == 400


def test_excel_follows_the_range(monkeypatch, funding) -> None:
    from openpyxl import load_workbook
    from app.api.insights import IST

    stub_model(monkeypatch, answer_all)
    rounds.extract_pending()
    start = str(utcnow().astimezone(IST).date() - timedelta(days=3))
    with TestClient(app) as client:
        resp = client.get("/api/insights/rounds.xlsx", params={"stage": "seed", "start": start})
        screen = client.get("/api/insights/rounds", params={"stage": "seed", "start": start}).json()
    ws = load_workbook(io.BytesIO(resp.content))["Seed"]
    assert ws.max_row - 1 == len(screen["rounds"])
    assert f"{start}-to-today" in resp.headers["content-disposition"]
