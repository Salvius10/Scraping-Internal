"""Insights "Search web" tests. Offline: Firecrawl and the model are stubbed.

Protected: the searches carry the reader's date range; only the four named
stages are kept; out-of-range, known and duplicate rounds are skipped; web
finds show in the tabs and Excel but never in the news feed; and a result
is read by the model only once.
"""

from __future__ import annotations

import io
from datetime import date, datetime, timedelta, timezone

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import delete, select

from app.config import settings
from app.db import session_scope
from app.ingest import web_rounds as wr
from app.llm.bedrock import LlmResult
from app.main import app
from app.models import Article, WebRound
from app.search import firecrawl as fc

IST = wr.IST
NOW = datetime(2026, 9, 25, 12, 0, tzinfo=timezone.utc)


@pytest.fixture(autouse=True)
def _clean(monkeypatch):
    monkeypatch.setattr(wr, "utcnow", lambda: NOW)
    yield
    with session_scope() as s:
        s.execute(delete(WebRound))


def result(n, url, title, published, snippet="") -> fc.SearchResult:
    return fc.SearchResult(n=n, url=url, title=title, description=snippet,
                           domain=fc.domain_of(url), published=published)


# One result set per stage query; the Qzorbit story appears under two queries.
RESULTS = {
    wr.Stage.PRE_SEED: [result(1, "https://a.test/qzorbit", "Qzorbit raises $1M pre-seed", "2 days ago")],
    wr.Stage.SEED: [
        result(1, "https://a.test/qzorbit", "Qzorbit raises $1M pre-seed", "2 days ago"),
        result(2, "https://b.test/zyro", "Zyro bags Rs 8 crore seed round", "Sep 20, 2026"),
    ],
    wr.Stage.SERIES_A: [result(1, "https://c.test/old", "Oldco Series A", "Aug 1, 2026")],
    wr.Stage.SERIES_B: [result(1, "https://d.test/amzn", "Amazon invests $3 Bn in unit", "1 day ago")],
}

MODEL = {
    "Qzorbit": {"company": "Qzorbit", "round": "pre seed", "stage": "Pre-Seed",
                "amount": "$1M", "investors": ["Acme VC"]},
    "Zyro": {"company": "Zyro", "round": "Seed", "stage": "Seed",
             "amount": "Rs 8 crore", "investors": []},
    "Amazon": {"company": "Amazon", "round": None, "stage": "Other",
               "amount": "$3 Bn", "investors": []},
}


def stub(monkeypatch, results=RESULTS):
    searches, prompts = [], []

    def fake_search(query, limit=None, kind="news", tbs=None, steered=True):
        searches.append({"query": query, "kind": kind, "tbs": tbs, "steered": steered})
        stage = next(s for s, q in wr.QUERIES.items() if q == query)
        return fc.Search(query=query, searched=query, kind=kind,
                         results=results[stage], credits_used=2)

    def fake_model(feature, prompt, **kw):
        prompts.append(prompt)
        out = []
        for line in prompt.splitlines():
            if line.startswith("[") and "]" in line:
                n, text = line[1:].split("]", 1)
                word = text.split()[0]
                if word in MODEL:
                    out.append({"i": int(n), **MODEL[word]})
        return out, LlmResult("", settings.model_cheap, 800, 200, 0.0002)

    monkeypatch.setattr(wr.firecrawl, "search", fake_search)
    monkeypatch.setattr(wr, "invoke_json", fake_model)
    return searches, prompts


def test_every_query_carries_the_date_range(monkeypatch) -> None:
    searches, _ = stub(monkeypatch)
    wr.search_web(date(2026, 9, 18), date(2026, 9, 25))
    assert len(searches) == 4
    assert all(s["tbs"] == "cdr:1,cd_min:9/18/2026,cd_max:9/25/2026" for s in searches)
    assert all(s["kind"] == "news" and s["steered"] is False for s in searches)


def test_only_new_in_range_startup_rounds_are_added(monkeypatch) -> None:
    stub(monkeypatch)
    summary = wr.search_web(date(2026, 9, 18), date(2026, 9, 25))

    assert summary.results_seen == 4                 # Qzorbit counted once
    assert summary.added == {"Pre-Seed": 1, "Seed": 1}
    assert summary.outside_range == 1                # Oldco, 1 Aug
    assert summary.not_startup_rounds == 1           # Amazon -> Other, dropped
    assert summary.credits_used == 8
    with session_scope() as s:
        rows = {w.company: w for w in s.scalars(
            select(WebRound).where(WebRound.kept.is_(True))).all()}
        hidden = s.scalars(select(WebRound.company).where(WebRound.kept.is_(False))).all()
    assert set(rows) == {"Qzorbit", "Zyro"}
    assert hidden == ["Amazon"]                      # recorded, never shown
    assert rows["Qzorbit"].round_label == "Pre-Seed"         # tidied
    assert rows["Qzorbit"].date_approx is True               # "2 days ago"
    assert rows["Zyro"].date_approx is False                 # a real date


def test_a_result_is_read_once(monkeypatch) -> None:
    _, prompts = stub(monkeypatch)
    wr.search_web(date(2026, 9, 18), date(2026, 9, 25))
    first = len(prompts)
    again = wr.search_web(date(2026, 9, 18), date(2026, 9, 25))
    assert len(prompts) == first                     # nothing new to read
    assert again.total_added == 0 and again.duplicates >= 2


def test_a_round_already_in_the_feed_is_skipped(monkeypatch) -> None:
    from app.ingest.rounds import Stage as S
    from app.models import Category, FundingRound

    stub(monkeypatch)
    with session_scope() as s:
        article = Article(url="https://feed.test/zyro", url_hash="z", source="test_web",
                          headline="Zyro raises seed", category=Category.FUNDING,
                          published_at=datetime(2026, 9, 21, tzinfo=timezone.utc))
        s.add(article)
        s.flush()
        s.add(FundingRound(article_id=article.id, company="zyro", stage=S.SEED))
    try:
        summary = wr.search_web(date(2026, 9, 18), date(2026, 9, 25))
        assert summary.added == {"Pre-Seed": 1}
    finally:
        with session_scope() as s:
            ids = s.scalars(select(Article.id).where(Article.source == "test_web")).all()
            s.execute(delete(FundingRound).where(FundingRound.article_id.in_(ids)))
            s.execute(delete(Article).where(Article.id.in_(ids)))


def test_partial_firecrawl_failure_is_reported_not_fatal(monkeypatch) -> None:
    stub(monkeypatch)
    real = wr.firecrawl.search

    def flaky(query, **kw):
        if query == wr.QUERIES[wr.Stage.SERIES_B]:
            raise fc.FirecrawlError("rate limit")
        return real(query, **kw)

    monkeypatch.setattr(wr.firecrawl, "search", flaky)
    summary = wr.search_web(date(2026, 9, 18), date(2026, 9, 25))
    assert summary.total_added == 2
    assert summary.failed_queries == ["Series B: rate limit"]


def test_total_failure_is_an_error(monkeypatch) -> None:
    def down(query, **kw):
        raise fc.FirecrawlError("Firecrawl credits are used up.")

    monkeypatch.setattr(wr.firecrawl, "search", down)
    with pytest.raises(wr.WebSearchError, match="credits are used up"):
        wr.search_web(date(2026, 9, 18))


def test_relative_and_absolute_dates() -> None:
    when, approx = wr.parse_published("3 days ago", NOW)
    assert approx and when.date() == date(2026, 9, 22)
    when, approx = wr.parse_published("Sep 20, 2026", NOW)
    assert not approx and when.astimezone(IST).date() == date(2026, 9, 20)
    assert wr.parse_published(None) == (None, False)


# --- API ---------------------------------------------------------------------------------

def test_web_finds_join_the_tabs_and_excel_but_not_the_feed(monkeypatch) -> None:
    from openpyxl import load_workbook

    stub(monkeypatch)
    with TestClient(app) as client:
        body = client.post("/api/insights/search-web",
                           json={"start": "2026-09-18", "end": "2026-09-25"}).json()
        seed = client.get("/api/insights/rounds", params={
            "stage": "seed", "start": "2026-09-18", "end": "2026-09-25"}).json()
        xlsx = client.get("/api/insights/rounds.xlsx", params={"stage": "seed"})
        feed = client.get("/api/articles", params={"q": "Zyro"}).json()

    assert body["error"] is None and body["total_added"] == 2
    zyro = next(r for r in seed["rounds"] if r["company"] == "Zyro")
    assert zyro["origin"] == "web" and zyro["source_label"] == "b.test"
    assert zyro["id"].startswith("web-")
    ws = load_workbook(io.BytesIO(xlsx.content))["Seed"]
    origins = [ws.cell(row=r, column=8).value for r in range(2, ws.max_row + 1)]
    assert "Web search" in origins
    assert all("Zyro" not in a["headline"] for a in feed["articles"])


def test_search_web_needs_a_start_date() -> None:
    with TestClient(app) as client:
        assert client.post("/api/insights/search-web", json={}).status_code == 422
