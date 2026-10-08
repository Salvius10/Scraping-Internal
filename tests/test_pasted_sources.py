"""Websites pasted on Insights > Startup firms and VC firms. Offline: the guarded
fetcher, Firecrawl and the model are stubbed.

Protected: a pasted link is read the cheapest way that finds posts (its feed,
a feed its page declares, its post links, then Firecrawl) and refused when it
is unusable or already read; each page keeps its own sources; a startup
website shows only the rounds in it, once, in the stage tabs; a VC website's
posts are sorted like a firm's own news and shown under the firm it was tied
to, else under its own name; post pages are opened through the guarded
fetcher; only new posts reach the model; a Firecrawl website is read at most
every `pasted_firecrawl_hours` and never with --skip-paid; removing a website
removes what was read from it; Excel marks where each row came from.
"""

from __future__ import annotations

import io
import re
from datetime import datetime, timedelta, timezone

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import delete, select

from app.config import settings
from app.db import session_scope
from app.extract import ExtractError
from app.ingest import pasted_sources as ps
from app.ingest import vc_firms as vf
from app.llm.bedrock import LlmResult
from app.main import app
from app.models import (
    Article, Category, FundingRound, PastedRound, PastedSource, PastedSourceRead, Stage,
    VcPost, utcnow,
)
from app.search import firecrawl as fc


def rss(*items: tuple[str, str]) -> str:
    body = "".join(
        f"<item><title>{title}</title><link>{url}</link>"
        f"<pubDate>Mon, 05 Oct 2026 10:00:00 +0530</pubDate>"
        f"<description>{title}.</description></item>" for title, url in items)
    return f'<?xml version="1.0"?><rss version="2.0"><channel><title>T</title>{body}</channel></rss>'


STARTUP_FEED = rss(
    ("Orbo raises $2M seed round led by Blume Ventures", "https://angels.test/orbo"),
    ("Kite launches a new app for small shops in Pune", "https://angels.test/kite"),
    ("Zeno raises $40M Series C to expand overseas", "https://angels.test/zeno"),
)
VC_FEED = rss(
    ("Backing Nimbus in its $3M seed round", "https://highway.test/nimbus"),
    ("Podcast: what we learnt about patience", "https://highway.test/podcast"),
)
NEWS_PAGE = """<html><body><nav><a href="/about">About the team and our story so far</a></nav>
<a href="/news/orbo">Orbo raises a $2M seed round from Blume</a>
<a href="/news/kite">Kite launches its new app for small shops</a>
<footer><a href="/privacy">Read our privacy policy and the terms</a></footer></body></html>"""
POST_PAGE = """<html><head><meta property="article:published_time" content="2026-10-04T09:30:00+05:30">
<meta name="description" content="The full story."></head><body></body></html>"""


@pytest.fixture(autouse=True)
def _clean():
    yield
    with session_scope() as s:
        s.execute(delete(PastedRound))
        s.execute(delete(PastedSourceRead))
        s.execute(delete(PastedSource))
        s.execute(delete(VcPost))
        s.execute(delete(FundingRound))
        s.execute(delete(Article).where(Article.url.like("https://feed.test/%")))


class Web(dict):
    """url -> page, and every url fetched."""

    def __init__(self) -> None:
        super().__init__()
        self.fetched: list[str] = []


@pytest.fixture
def web(monkeypatch) -> Web:
    """The internet as a dict of url -> page. Anything else answers 404. The
    model and Firecrawl fail the test unless a test stubs them."""
    pages = Web()

    def fetch(url):
        pages.fetched.append(url)
        if url not in pages:
            raise ExtractError("The site answered 404 for that link.")
        return url, pages[url]

    monkeypatch.setattr(ps, "check_url", lambda url: url)
    monkeypatch.setattr(ps, "fetch_page", fetch)
    monkeypatch.setattr(ps.firecrawl, "scrape", lambda url: pytest.fail("Firecrawl was called"))
    monkeypatch.setattr(ps, "invoke_json", lambda *a, **k: pytest.fail("the model was called"))
    monkeypatch.setattr(vf, "invoke_json", lambda *a, **k: pytest.fail("the model was called"))
    monkeypatch.setattr(settings, "per_domain_delay", 0)
    return pages


def stub_rounds(monkeypatch, answers: dict[str, dict]) -> list[str]:
    """The model answers by the first word of each numbered post."""
    prompts = []

    def fake(feature, prompt, **kw):
        prompts.append(prompt)
        out = []
        for line in prompt.splitlines():
            m = re.match(r"\[(\d+)\] (\S+)", line)
            if m and m.group(2) in answers:
                out.append({"i": int(m.group(1)), **answers[m.group(2)]})
        return out, LlmResult("", settings.model_cheap, 900, 150, 0.0003)

    monkeypatch.setattr(ps, "invoke_json", fake)
    return prompts


def stub_sort(monkeypatch, answers: dict[str, dict]) -> list[str]:
    prompts = []

    def fake(feature, prompt, **kw):
        prompts.append(prompt)
        out = []
        for line in prompt.splitlines():
            m = re.match(r"\[(\d+)\] \(Firm: [^)]+\) (\S+)", line)
            if m and m.group(2) in answers:
                out.append({"i": int(m.group(1)), **answers[m.group(2)]})
        return out, LlmResult("", settings.model_cheap, 900, 150, 0.0002)

    monkeypatch.setattr(vf, "invoke_json", fake)
    return prompts


ROUNDS = {
    "Orbo": {"company": "Orbo", "round": "seed", "stage": "Seed", "amount": "$2M",
             "investors": ["Blume Ventures"], "raise": True},
    "Kite": {"company": "Kite", "round": None, "stage": "Other", "amount": None,
             "investors": [], "raise": False},
    "Zeno": {"company": "Zeno", "round": "Series C", "stage": "Other", "amount": "$40M",
             "investors": [], "raise": True},
}
SORTS = {
    "Backing": {"kind": "Investment", "company": "Nimbus", "round": "seed", "amount": "$3M",
                "headline": "Venture Highway backs Nimbus in $3M seed"},
    "Podcast:": {"kind": "Other", "company": None, "headline": "On patience"},
}


# --- How a link is read ---------------------------------------------------------------

def test_a_feed_link_is_read_as_a_feed(web) -> None:
    web["https://angels.test/feed"] = STARTUP_FEED
    assert ps.plan_read("https://angels.test/feed") == ps.Plan("rss", "https://angels.test/feed")


def test_a_page_that_declares_a_feed_is_read_from_it(web) -> None:
    web["https://angels.test/news"] = (
        '<html><head><link rel="alternate" type="application/rss+xml" title="Comments Feed" '
        'href="/comments/feed"><link rel="alternate" type="application/rss+xml" '
        'href="/feed"></head><body></body></html>')
    web["https://angels.test/feed"] = STARTUP_FEED
    assert ps.plan_read("https://angels.test/news") == ps.Plan("rss", "https://angels.test/feed")
    assert "https://angels.test/comments/feed" not in web.fetched


def test_a_page_with_its_own_feed_is_read_from_that_one(web) -> None:
    """A WordPress category page lists the whole site's feed first."""
    web["https://site.test/category/venture/"] = (
        '<html><head><link rel="alternate" type="application/rss+xml" href="/feed/">'
        '<link rel="alternate" type="application/rss+xml" '
        'href="https://site.test/category/venture/feed/"></head></html>')
    web["https://site.test/category/venture/feed/"] = VC_FEED
    assert ps.plan_read("https://site.test/category/venture/") == ps.Plan(
        "rss", "https://site.test/category/venture/feed/")


def test_a_feed_declared_in_another_encoding_reads_correctly(web) -> None:
    # fetch_page hands over decoded text; the declaration must not re-decode it.
    web["https://site.test/feed"] = rss(("Café Coffee raises ₹40 crore in a Series B",
                                         "https://site.test/cafe")).replace(
        '<?xml version="1.0"?>', '<?xml version="1.0" encoding="ISO-8859-1"?>')
    found, _ = ps._fetch(PastedSource(via="rss", read_url="https://site.test/feed"))
    assert found[0].title == "Café Coffee raises ₹40 crore in a Series B"


def test_a_news_page_is_read_for_its_post_links(web) -> None:
    web["https://angels.test/news"] = NEWS_PAGE
    assert ps.plan_read("https://angels.test/news").via == "html"


def test_a_page_that_refuses_plain_http_falls_back_to_firecrawl(web, monkeypatch) -> None:
    monkeypatch.setattr(settings, "firecrawl_api_key", "fc-test")
    assert ps.plan_read("https://blocked.test/news") == ps.Plan("scrape", "https://blocked.test/news")

    monkeypatch.setattr(settings, "firecrawl_api_key", None)
    monkeypatch.setattr(settings, "firecrawl_api_url", "https://api.firecrawl.dev")
    with pytest.raises(ps.SourceError, match="FIRECRAWL_API_KEY"):
        ps.plan_read("https://blocked.test/news")


@pytest.mark.parametrize("body, words", [
    ({"section": "startups", "url": ""}, "Paste a link"),
    ({"section": "startups", "url": "not a link"}, "does not look like"),
    ({"section": "everything", "url": "https://angels.test/feed"}, "Unknown section"),
    ({"section": "vcs", "url": "https://angels.test/feed", "firm": "nobody"}, "Unknown firm"),
    ({"section": "startups", "url": "https://entrackr.com/rss"}, "already in the news feed"),
])
def test_add_refuses_what_it_cannot_use(web, body, words) -> None:
    out = TestClient(app).post("/api/insights/sources", json=body).json()
    assert not out["added"] and words in out["error"]


def test_a_vc_firms_own_news_page_is_not_added_twice(web) -> None:
    accel = vf.get_firm("accel")
    out = TestClient(app).post("/api/insights/sources",
                               json={"section": "vcs", "url": accel.site.url}).json()
    assert not out["added"] and "already read from that link" in out["error"]


def test_each_page_keeps_its_own_sources(web, monkeypatch) -> None:
    web["https://angels.test/feed"] = STARTUP_FEED
    stub_rounds(monkeypatch, ROUNDS)
    client = TestClient(app)
    body = {"section": "startups", "url": "https://angels.test/feed"}
    assert client.post("/api/insights/sources", json=body).json()["added"]
    again = client.post("/api/insights/sources", json=body).json()
    assert not again["added"] and "already a source here" in again["error"]

    stub_sort(monkeypatch, {})
    assert client.post("/api/insights/sources", json={**body, "section": "vcs"}).json()["added"]
    assert len(client.get("/api/insights/rounds").json()["sources"]) == 1
    assert len(client.get("/api/insights/vcs").json()["sources"]) == 1


# --- Startup firms -----------------------------------------------------------------------

def test_a_startup_website_shows_only_its_rounds(web, monkeypatch) -> None:
    web["https://angels.test/feed"] = STARTUP_FEED
    prompts = stub_rounds(monkeypatch, ROUNDS)
    client = TestClient(app)
    out = client.post("/api/insights/sources", json={
        "section": "startups", "url": "https://angels.test/feed", "label": "Angels Weekly"}).json()

    assert out["added"] and out["via"] == "rss" and out["error"] is None
    assert (out["found"], out["new"], out["shown"], out["credits_used"]) == (3, 3, 2, 0)
    assert ps.PASTED_NOTE in prompts[0]

    seed = client.get("/api/insights/rounds", params={"stage": "seed"}).json()
    row = next(r for r in seed["rounds"] if r["origin"] == "pasted")
    assert (row["company"], row["round"], row["amount"], row["investors"], row["source_label"]) == (
        "Orbo", "Seed", "$2M", ["Blume Ventures"], "Angels Weekly")
    other = client.get("/api/insights/rounds", params={"stage": "other"}).json()
    assert [r["company"] for r in other["rounds"] if r["origin"] == "pasted"] == ["Zeno"]
    # A launch is not a round: read once, never shown.
    assert "Kite" not in str(seed) + str(other)
    assert {s["key"]: s["count"] for s in seed["stages"]}["seed"] >= 1
    assert seed["sources"][0]["shown"] == 2 and seed["sources"][0]["via_label"] == "its feed"


def test_a_round_already_in_the_feed_shows_once(web, monkeypatch) -> None:
    with session_scope() as s:
        a = Article(url="https://feed.test/orbo", url_hash="psorbo", source="entrackr",
                    headline="Orbo raises $2M", category=Category.FUNDING,
                    published_at=datetime(2026, 10, 4, tzinfo=timezone.utc))
        s.add(a)
        s.flush()
        s.add(FundingRound(article_id=a.id, company="Orbo", stage=Stage.SEED, amount="$2M"))
    web["https://angels.test/feed"] = STARTUP_FEED
    stub_rounds(monkeypatch, ROUNDS)
    out = TestClient(app).post("/api/insights/sources",
                               json={"section": "startups", "url": "https://angels.test/feed"}).json()
    assert out["shown"] == 1                                   # only Zeno is new
    with session_scope() as s:
        orbo = s.scalar(select(PastedRound).where(PastedRound.company == "Orbo"))
    assert orbo.extracted_at is not None and orbo.kept is False


def test_only_new_posts_reach_the_model(web, monkeypatch) -> None:
    web["https://angels.test/feed"] = STARTUP_FEED
    prompts = stub_rounds(monkeypatch, ROUNDS)
    source = ps.add_source("startups", "https://angels.test/feed")
    ps.read_now(source)
    assert len(prompts) == 1

    result = ps.refresh_pasted()
    assert result.reads[0].items_new == 0 and len(prompts) == 1


def test_news_page_posts_get_their_dates_from_their_own_pages(web, monkeypatch) -> None:
    web["https://angels.test/news"] = NEWS_PAGE
    web["https://angels.test/news/orbo"] = POST_PAGE
    stub_rounds(monkeypatch, ROUNDS)
    source = ps.add_source("startups", "https://angels.test/news")
    run = ps.read_source(source)
    assert (run.items_seen, run.items_new, run.error) == (2, 2, None)
    with session_scope() as s:
        rows = {r.url: r for r in s.scalars(select(PastedRound)).all()}
    orbo = rows["https://angels.test/news/orbo"]
    assert orbo.published_at.isoformat() == "2026-10-04T04:00:00+00:00"
    assert orbo.date_approx is False and orbo.snippet == "The full story"
    # Opened through the guarded fetcher, like the page itself.
    assert "https://angels.test/news/kite" in web.fetched


def test_a_website_with_no_posts_is_reported(web) -> None:
    web["https://angels.test/news"] = NEWS_PAGE
    source = ps.add_source("startups", "https://angels.test/news")
    web["https://angels.test/news"] = "<html><body>Nothing here</body></html>"
    run = ps.read_source(source)
    assert run.error and "No posts found" in run.error


# --- VC firms ------------------------------------------------------------------------------

def test_a_vc_website_is_sorted_like_firm_news(web, monkeypatch) -> None:
    web["https://highway.test/feed"] = VC_FEED
    prompts = stub_sort(monkeypatch, SORTS)
    client = TestClient(app)
    out = client.post("/api/insights/sources", json={
        "section": "vcs", "url": "https://highway.test/feed", "label": "Venture Highway"}).json()
    assert out["added"] and (out["found"], out["new"], out["shown"]) == (2, 2, 1)
    assert "(Firm: Venture Highway)" in prompts[0]

    key = f"pasted-{out['source_id']}"
    data = client.get("/api/insights/vcs").json()
    post = next(p for p in data["posts"] if p["origin"] == "added")
    assert (post["firm"], post["firm_label"], post["kind"], post["source_label"]) == (
        key, "Venture Highway", "Investment", "highway.test")
    assert data["hidden_other"] >= 1                          # the podcast
    firm = next(f for f in data["firms"] if f["key"] == key)
    assert firm["pasted"] and firm["count"] == 1
    assert firm["reads"][0]["label"] == "added by you, from its feed"

    assert len(client.get("/api/insights/vcs", params={"firm": key}).json()["posts"]) == 1
    only = client.get("/api/insights/vcs", params={"origin": "added"}).json()["posts"]
    assert [p["company"] for p in only] == ["Nimbus"]


def test_a_vc_website_tied_to_a_firm_shows_under_it(web, monkeypatch) -> None:
    web["https://highway.test/feed"] = VC_FEED
    stub_sort(monkeypatch, SORTS)
    client = TestClient(app)
    out = client.post("/api/insights/sources", json={
        "section": "vcs", "url": "https://highway.test/feed", "firm": "accel"}).json()
    assert out["added"]

    data = client.get("/api/insights/vcs", params={"firm": "accel"}).json()
    added = [p for p in data["posts"] if p["origin"] == "added"]
    assert [(p["firm"], p["company"]) for p in added] == [("accel", "Nimbus")]
    accel = next(f for f in data["firms"] if f["key"] == "accel")
    assert any(r["label"].startswith("added by you") for r in accel["reads"])
    assert not any(f["key"] == f"pasted-{out['source_id']}" for f in data["firms"])
    assert data["sources"][0]["firm_label"] == accel["label"]


# --- Refresh ---------------------------------------------------------------------------------

def test_a_firecrawl_website_is_read_at_most_daily(web, monkeypatch) -> None:
    monkeypatch.setattr(settings, "firecrawl_api_key", "fc-test")
    scrapes = []

    def scrape(url):
        scrapes.append(url)
        return fc.Page(url=url, title="News", description="", credits_used=1,
                       markdown="[Orbo raises a $2M seed round from Blume](https://js.test/orbo)"), False

    monkeypatch.setattr(ps.firecrawl, "scrape", scrape)
    source = ps.add_source("startups", "https://js.test/news")
    assert source.via == "scrape" and scrapes == []           # adding costs nothing

    first = ps.refresh_pasted(classify=False)
    assert first.credits == 1 and first.new == 1
    second = ps.refresh_pasted(classify=False)
    assert second.reads == [] and second.skipped == 1 and len(scrapes) == 1

    with session_scope() as s:
        for r in s.scalars(select(PastedSourceRead)).all():
            r.started_at -= timedelta(hours=settings.pasted_firecrawl_hours)
    assert ps.refresh_pasted(free_only=True, classify=False).skipped == 1
    assert len(ps.refresh_pasted(classify=False).reads) == 1


def test_skip_paid_still_reads_free_websites(web) -> None:
    web["https://angels.test/feed"] = STARTUP_FEED
    ps.add_source("startups", "https://angels.test/feed")
    result = ps.refresh_pasted(free_only=True, classify=False)     # the model is never called
    assert result.new == 3 and result.rounds is None


# --- Remove, Excel ---------------------------------------------------------------------------

def test_removing_a_website_removes_what_was_read(web, monkeypatch) -> None:
    web["https://angels.test/feed"] = STARTUP_FEED
    web["https://highway.test/feed"] = VC_FEED
    stub_rounds(monkeypatch, ROUNDS)
    stub_sort(monkeypatch, SORTS)
    client = TestClient(app)
    ids = [client.post("/api/insights/sources", json=b).json()["source_id"] for b in (
        {"section": "startups", "url": "https://angels.test/feed"},
        {"section": "vcs", "url": "https://highway.test/feed"},
    )]
    for source_id in ids:
        assert client.delete(f"/api/insights/sources/{source_id}").json() == {"removed": True}
    assert client.delete(f"/api/insights/sources/{ids[0]}").status_code == 404

    with session_scope() as s:
        assert s.scalars(select(PastedRound)).all() == []
        assert s.scalars(select(VcPost).where(VcPost.via == "added")).all() == []
        assert s.scalars(select(PastedSourceRead)).all() == []


def test_excel_says_where_each_row_came_from(web, monkeypatch) -> None:
    from openpyxl import load_workbook

    web["https://angels.test/feed"] = STARTUP_FEED
    web["https://highway.test/feed"] = VC_FEED
    stub_rounds(monkeypatch, ROUNDS)
    stub_sort(monkeypatch, SORTS)
    client = TestClient(app)
    client.post("/api/insights/sources", json={"section": "startups", "url": "https://angels.test/feed"})
    client.post("/api/insights/sources", json={"section": "vcs", "url": "https://highway.test/feed"})

    rounds = load_workbook(io.BytesIO(
        client.get("/api/insights/rounds.xlsx", params={"stage": "seed"}).content)).active
    row = next(r for r in rounds.iter_rows(values_only=True) if r[0] == "Orbo")
    assert row[6] == "angels.test" and row[7] == "Your website"

    vcs = load_workbook(io.BytesIO(client.get("/api/insights/vcs.xlsx").content)).active
    row = next(r for r in vcs.iter_rows(values_only=True) if r[2] == "Nimbus")
    assert row[0] == "highway.test" and row[6] == "Added by you: highway.test"
