"""Platform endpoints: the two feeds, Export to Excel, status, refresh, the log.

Protected: filters and counts agree with the rows; upcoming and past split
every event between them; the Excel file holds what the page shows, with a
clickable source URL and no cell that runs as a formula; status reports
freshness, guardrails and every source; a refused refresh says why.
"""

from __future__ import annotations

import io
import threading
import time
from datetime import datetime, timedelta, timezone

import pytest
from fastapi.testclient import TestClient

from app import scheduler
from app.db import session_scope
from app.main import app
from app.models import BlockedItem, FundedStartup, PipelineRun, Stage, StartupEvent, utcnow

client = TestClient(app)
IST = timezone(timedelta(hours=5, minutes=30))


def funded(company: str, stage: Stage, published: datetime | None, **over) -> FundedStartup:
    values = dict(company=company, company_key=company.lower(), stage=stage, round_label=None,
                  amount="$1M", investors=["Fund One"], sector="Fintech", hq_city="Pune",
                  summary=f"{company} raised.", headline=f"{company} raises",
                  source_url=f"https://news.test/{company.lower()}", source_label="News",
                  source_type="rss", published_at=published,
                  sources=[{"url": f"https://news.test/{company.lower()}", "source": "News",
                            "type": "rss", "published_at": None}])
    values.update(over)
    return FundedStartup(**values)


def event(title: str, starts: datetime, **over) -> StartupEvent:
    values = dict(title=title, title_key=title.lower(), starts_at=starts, ends_at=None,
                  date_only=False, city="Bengaluru", url=f"https://lu.ma/{title[:3]}",
                  source_url=f"https://lu.ma/{title[:3]}", source_label="Luma",
                  source_type="luma", sources=[], category="Meetup")
    values.update(over)
    return StartupEvent(**values)


@pytest.fixture
def data(clean_db):
    now = utcnow()
    with session_scope() as s:
        s.add_all([
            funded("Acme", Stage.SEED, datetime(2026, 10, 10, 20, 0, tzinfo=timezone.utc)),  # 11 Oct IST
            funded("Beta", Stage.SERIES_A, datetime(2026, 10, 5, 6, 0, tzinfo=timezone.utc),
                   source_type="linkedin"),
            funded("=HYPERLINK(\"http://evil\")", Stage.SEED, None, sector="100%"),
            event("Upcoming Meetup", now + timedelta(days=2)),
            event("Running Summit", now - timedelta(hours=2), ends_at=now + timedelta(hours=3),
                  city="Mumbai"),
            event("Old Demo Day", now - timedelta(days=3)),
            event("Today Only", now - timedelta(hours=1), date_only=True, city=None),
        ])


def test_funding_filters_and_counts(data) -> None:
    body = client.get("/api/funding").json()
    assert body["total"] == 3
    assert [r["company"] for r in body["rows"]][:2] == ["Acme", "Beta"]     # newest first
    counts = {s["key"]: s["count"] for s in body["stages"]}
    assert counts == {"pre-seed": 0, "seed": 2, "series-a": 1, "series-b": 0, "other": 0}

    seed = client.get("/api/funding", params={"stage": "seed"}).json()
    assert {r["company"] for r in seed["rows"]} == {"Acme", "=HYPERLINK(\"http://evil\")"}
    assert {s["key"]: s["count"] for s in seed["stages"]}["series-a"] == 1   # counts ignore the tab

    dated = client.get("/api/funding", params={"start": "2026-10-11", "end": "2026-10-11"}).json()
    assert [r["company"] for r in dated["rows"]] == ["Acme"]                 # IST day, not UTC
    assert client.get("/api/funding", params={"source": "linkedin"}).json()["total"] == 1
    assert client.get("/api/funding", params={"q": "100%"}).json()["total"] == 1
    assert client.get("/api/funding", params={"q": "zzz"}).json()["total"] == 0


def test_funding_refuses_bad_filters(data) -> None:
    assert client.get("/api/funding", params={"stage": "series-z"}).status_code == 400
    assert client.get("/api/funding", params={"source": "x"}).status_code == 400
    assert client.get("/api/funding", params={"start": "2026-10-11", "end": "2026-10-01"}
                      ).status_code == 400
    assert client.get("/api/funding", params={"limit": 100000}).status_code == 422


def test_every_event_is_upcoming_or_past(data) -> None:
    upcoming = client.get("/api/events").json()
    past = client.get("/api/events", params={"when": "past"}).json()
    assert {r["title"] for r in upcoming["rows"]} == {"Upcoming Meetup", "Running Summit",
                                                      "Today Only"}
    assert [r["title"] for r in past["rows"]] == ["Old Demo Day"]
    assert upcoming["upcoming"] == 3 and upcoming["past"] == 1
    starts = [r["starts_at"] for r in upcoming["rows"]]
    assert starts == sorted(starts)                                          # soonest first
    assert client.get("/api/events", params={"when": "all"}).json()["total"] == 4
    cities = {c["name"]: c["count"] for c in upcoming["cities"]}
    assert cities == {"Bengaluru": 1, "Mumbai": 1}
    assert client.get("/api/events", params={"city": "mumbai"}).json()["total"] == 1
    assert client.get("/api/events", params={"when": "soon"}).status_code == 400


def test_funding_excel_holds_the_page_and_runs_no_formula(data) -> None:
    from openpyxl import load_workbook

    response = client.get("/api/export/funding.xlsx", params={"stage": "seed"})
    assert response.status_code == 200
    assert "attachment" in response.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(response.content)).active
    header = [c.value for c in ws[1]]
    assert header[0] == "Company" and "Source URL" in header and "Published (IST)" in header
    rows = list(ws.iter_rows(min_row=2))
    assert len(rows) == 2
    evil = next(r for r in rows if str(r[0].value).startswith("=HYPERLINK"))
    assert evil[0].data_type == "s"                                         # text, not a formula
    acme = next(r for r in rows if r[0].value == "Acme")
    link = acme[header.index("Source URL")]
    assert link.value == "https://news.test/acme" and link.hyperlink.target == link.value
    assert acme[header.index("Published (IST)")].value == datetime(2026, 10, 11, 1, 30)


def test_events_excel_follows_the_view(data) -> None:
    from openpyxl import load_workbook

    response = client.get("/api/export/events.xlsx", params={"when": "past"})
    ws = load_workbook(io.BytesIO(response.content)).active
    assert [r[0].value for r in ws.iter_rows(min_row=2)] == ["Old Demo Day"]


def test_status_reports_freshness_guardrails_and_sources(data) -> None:
    with session_scope() as s:
        s.add(PipelineRun(trigger="scheduled", started_at=utcnow() - timedelta(hours=3),
                          finished_at=utcnow() - timedelta(hours=2), fetched=10, admitted=4,
                          funding_new=2, events_new=1, blocked_raw=1))
        s.add(BlockedItem(stage="guardrail_raw", source_key="rss:a", url="https://x",
                          title="x", reason="content filter: HATE (HIGH)"))
    body = client.get("/api/status").json()
    assert body["hours_since_refresh"] == pytest.approx(2.0, abs=0.1)
    assert body["last_run"]["funding_new"] == 2 and body["last_run"]["blocked"] == 1
    assert body["guardrails_on"] is False and body["guardrails"].startswith("off")
    assert body["funding_count"] == 3 and body["upcoming_events"] == 3 and body["blocked"] == 1
    assert {s["key"] for s in body["sources"]} >= {"rss:entrackr", "luma", "eventbrite"}
    assert body["spend"]["total_cap"] == 7.0 and body["registry_error"] is None


def test_blocked_log_is_readable(data) -> None:
    with session_scope() as s:
        s.add(BlockedItem(stage="validation", source_key="rss:a", url="https://y", title="y",
                          reason="the model found a round but named no company"))
    rows = client.get("/api/blocked", params={"stage": "validation"}).json()
    assert [r["url"] for r in rows] == ["https://y"]
    assert client.get("/api/blocked", params={"stage": "nope"}).status_code == 400


# --- Refresh --------------------------------------------------------------------------

def wait_until_done(timeout: float = 5.0) -> scheduler.ManualRun:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        state = scheduler.manual_status()
        if not state.running:
            return state
        time.sleep(0.02)
    raise AssertionError("manual refresh did not finish")


def test_refresh_runs_in_the_background_and_reports(monkeypatch, clean_db) -> None:
    gate = threading.Event()
    monkeypatch.setattr(scheduler, "last_refresh", lambda: None)

    def fake_run(trigger):
        assert trigger == "manual"
        gate.wait(5)
        with session_scope() as s:
            run = PipelineRun(trigger=trigger, fetched=5, admitted=2, funding_new=1,
                              events_new=1, finished_at=utcnow())
            s.add(run)
            s.flush()
            return run

    monkeypatch.setattr(scheduler, "_run", fake_run)
    first = client.post("/api/refresh").json()
    assert first["started"] is True and first["status"]["running"] is True
    second = client.post("/api/refresh").json()
    assert second["started"] is False and "already running" in second["reason"]
    gate.set()
    done = wait_until_done()
    assert done.error is None and done.summary.startswith("Refreshed: 1 new funded startup")
    assert client.get("/api/refresh").json()["running"] is False


def test_refresh_cooldown_refuses_a_second_paid_run(monkeypatch, clean_db) -> None:
    monkeypatch.setattr(scheduler, "last_refresh", lambda: utcnow() - timedelta(minutes=2))
    answer = client.post("/api/refresh").json()
    assert answer["started"] is False and "Try again in 8 minutes" in answer["reason"]
