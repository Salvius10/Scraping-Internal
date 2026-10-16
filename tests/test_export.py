"""Export tests. Offline: rows are written straight into the test database.

Protected: company names match in one spelling; one deal from two sources
becomes one row, the trusted copy kept and filled from the other, Found via
naming both; different stages, far-apart dates, undated rows and unnamed
companies are never merged; a selection is a saved filter, so a round that
arrives later appears without adding it again; adding reports how many rows
were new and how many were duplicates; the same selection cannot be added
twice and bad ones are refused in words; LinkedIn adds only Funding posts;
selections can be taken back or cleared, and go with a removed LinkedIn
account; the Excel file is the Startup firms layout, a sheet per stage.

Company names here are unusual on purpose: the test database is shared, and
the merge must not meet rows other test files leave behind.
"""

from __future__ import annotations

import io
from datetime import datetime, timedelta, timezone

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import delete

from app.api.export import company_key, merge
from app.api.insights import COLUMNS, STAGE_TITLES, RoundOut
from app.db import session_scope
from app.ingest import linkedin as li
from app.main import app
from app.models import (
    Article, Category, ExportSelection, FundingRound, LinkedinKind, LinkedinPost,
    LinkedinRead, LinkedinSource, Stage, utcnow,
)

SOURCE = "test_export"
client = TestClient(app)


@pytest.fixture(autouse=True)
def _clean():
    yield
    with session_scope() as s:
        s.execute(delete(ExportSelection))
        s.execute(delete(LinkedinPost))
        s.execute(delete(LinkedinRead))
        s.execute(delete(LinkedinSource))
        ids = [a.id for a in s.query(Article).filter(Article.source == SOURCE)]
        s.execute(delete(FundingRound).where(FundingRound.article_id.in_(ids)))
        s.execute(delete(Article).where(Article.id.in_(ids)))


def feed_round(company: str, stage: Stage, days_ago: float, *, amount: str | None = None,
               investors: str | None = None, round_label: str | None = None) -> int:
    """A canonical Funding story in the feed, with its extracted round."""
    with session_scope() as s:
        n = s.query(Article).count()
        article = Article(url=f"https://export.test/{company}/{n}", url_hash=f"x{n}",
                          source=SOURCE, headline=f"{company} raises money",
                          category=Category.FUNDING,
                          published_at=utcnow() - timedelta(days=days_ago))
        s.add(article)
        s.flush()
        s.add(FundingRound(article_id=article.id, company=company, stage=stage,
                           round_label=round_label or STAGE_TITLES[stage], amount=amount,
                           investors=investors))
        return article.id


def linkedin_account(label: str = "Acme Ventures") -> int:
    with session_scope() as s:
        src = LinkedinSource(url=f"https://www.linkedin.com/company/{label.lower().replace(' ', '-')}/",
                             kind="company", label=label)
        s.add(src)
        s.flush()
        return src.id


def linkedin_post(source_id: int, kind: LinkedinKind, company: str, days_ago: float, *,
                  round_label: str | None = None, amount: str | None = None,
                  investors: str | None = None) -> None:
    with session_scope() as s:
        n = s.query(LinkedinPost).count()
        s.add(LinkedinPost(
            source_id=source_id, post_id=f"p{n}", url=f"https://www.linkedin.com/posts/x_{n}",
            content=f"{company} news", posted_at=utcnow() - timedelta(days=days_ago),
            classified_at=utcnow(), kind=kind, headline=f"{company} on LinkedIn",
            company=company, round_label=round_label, amount=amount, investors=investors))


def add(**body) -> dict:
    return client.post("/api/export/selections", json=body).json()


def ours(body: dict) -> dict[str, dict]:
    """Our rows in the Export list, by company."""
    names = {"Orbotron", "Zetaplex", "Kitewharf", "Lumenquay"}
    return {r["company"]: r for r in body["rounds"] if r["company"] in names}


# --- Merging ------------------------------------------------------------------------

@pytest.mark.parametrize("a,b", [
    ("Orbo", "orbo"), ("Orbo Pvt. Ltd.", "Orbo"), ("Orbo Private Limited", "ORBO"),
    ("Kite-Wharf", "kite wharf"),
])
def test_company_names_match_in_one_spelling(a, b) -> None:
    assert company_key(a) == company_key(b)


def test_company_key_keeps_names_that_are_only_a_suffix() -> None:
    assert company_key("Limited") == "limited"
    assert company_key("") is None and company_key(None) is None


def row(id: str, origin: str, company: str | None, stage: str, when: datetime | None,
        **extra) -> RoundOut:
    return RoundOut(id=id, company=company, stage=stage, round=extra.pop("round", None),
                    amount=extra.pop("amount", None), investors=extra.pop("investors", []),
                    published_at=when, source=origin, source_label=origin, headline=id,
                    url=f"https://x.test/{id}", origin=origin, **extra)


NOW = datetime(2026, 10, 10, 6, 0, tzinfo=timezone.utc)


def test_one_deal_from_two_sources_is_one_row_filled_from_both() -> None:
    merged = merge([
        row("li-1", "linkedin", "Orbo", "Seed", NOW, amount="$2M",
            investors=["Blume Ventures"], round="Seed"),
        row("feed-1", "feed", "Orbo Pvt Ltd", "Seed", NOW - timedelta(days=3)),
    ])
    assert len(merged) == 1
    [kept] = merged
    assert kept.id == "feed-1" and kept.also_via == ["linkedin"]     # the news copy is kept
    assert (kept.amount, kept.investors, kept.round) == ("$2M", ["Blume Ventures"], "Seed")


def test_what_is_never_merged() -> None:
    rows = [
        row("a", "feed", "Orbo", "Seed", NOW),
        row("b", "linkedin", "Orbo", "Series A", NOW),                    # other stage
        row("c", "web", "Orbo", "Seed", NOW - timedelta(days=30)),        # weeks apart
        row("d", "pasted", "Orbo", "Seed", None),                         # undated
        row("e", "feed", None, "Seed", NOW),                              # no company
        row("f", "linkedin", None, "Seed", NOW),
    ]
    assert len(merge(rows)) == 6


def test_the_same_row_twice_counts_once_and_order_is_newest_first() -> None:
    rows = [row("a", "feed", "Orbo", "Seed", NOW - timedelta(days=40)),
            row("b", "web", "Kite", "Seed", NOW),
            row("a", "feed", "Orbo", "Seed", NOW - timedelta(days=40)),
            row("c", "web", "Lumen", "Seed", None)]
    assert [r.id for r in merge(rows)] == ["b", "a", "c"]


# --- API ------------------------------------------------------------------------------

def test_adding_reports_new_rows_and_duplicates() -> None:
    feed_round("Orbotron", Stage.SEED, 1, round_label="Seed")
    feed_round("Kitewharf", Stage.SERIES_A, 3, amount="$5M", investors="Accel")
    acme = linkedin_account()
    linkedin_post(acme, LinkedinKind.FUNDING, "Orbotron", 2, round_label="Seed",
                  amount="$2M", investors="Acme Ventures; Blume Ventures")
    linkedin_post(acme, LinkedinKind.FUNDING, "Zetaplex", 4, round_label="Pre-Series A")
    linkedin_post(acme, LinkedinKind.NEWS, "Lumenquay", 1)               # never exported

    first = add(section="startups")
    assert first["added"] and first["label"] == "Startup firms · all stages"
    assert first["duplicates"] == 0 and first["new"] == first["rounds"] >= 2

    second = add(section="linkedin")
    assert second["added"] and second["label"] == "LinkedIn · all accounts"
    # Two Funding posts: Orbotron is already there from the news, Zetaplex is new.
    assert (second["rounds"], second["new"], second["duplicates"]) == (2, 1, 1)
    assert second["total"] == first["total"] + 1

    body = client.get("/api/export").json()
    mine = ours(body)
    assert set(mine) == {"Orbotron", "Kitewharf", "Zetaplex"}             # no News post
    orbo = mine["Orbotron"]
    assert orbo["origin"] == "feed" and orbo["also_via"] == ["linkedin"]
    assert orbo["amount"] == "$2M" and orbo["investors"] == ["Acme Ventures", "Blume Ventures"]
    assert mine["Zetaplex"]["stage"] == "Seed"                           # Pre-Series A is Seed
    assert body["duplicates"] >= 1 and body["total"] == len(body["rounds"])
    assert [s["label"] for s in body["selections"]] == ["Startup firms · all stages",
                                                        "LinkedIn · all accounts"]
    assert body["selections"][1]["rounds"] == 2


def test_a_selection_stays_fresh() -> None:
    add(section="startups", stage="seed")
    assert "Orbotron" not in ours(client.get("/api/export").json())
    feed_round("Orbotron", Stage.SEED, 0.5)                              # arrives later
    feed_round("Kitewharf", Stage.SERIES_A, 0.5)                         # not in the selection
    mine = ours(client.get("/api/export").json())
    assert set(mine) == {"Orbotron"}


def test_dates_and_one_stage_or_account_are_kept() -> None:
    feed_round("Orbotron", Stage.SEED, 1)
    feed_round("Kitewharf", Stage.SEED, 20)
    start = (utcnow() - timedelta(days=5)).astimezone(timezone(timedelta(hours=5, minutes=30)))
    result = add(section="startups", stage="seed", start=start.date().isoformat())
    assert result["label"] == "Startup firms · Seed"
    body = client.get("/api/export").json()
    assert set(ours(body)) == {"Orbotron"}
    assert body["selections"][0]["dates"].endswith("to today")

    one, other = linkedin_account("Acme Ventures"), linkedin_account("Beta Capital")
    linkedin_post(one, LinkedinKind.FUNDING, "Zetaplex", 1, round_label="Seed")
    linkedin_post(other, LinkedinKind.FUNDING, "Lumenquay", 1, round_label="Seed")
    assert add(section="linkedin", source_id=one)["label"] == "LinkedIn · Acme Ventures"
    assert set(ours(client.get("/api/export").json())) == {"Orbotron", "Zetaplex"}


@pytest.mark.parametrize("body,says", [
    ({"section": "vcs"}, "Only Startup firms and LinkedIn"),
    ({"section": "startups", "stage": "series-z"}, "Unknown stage"),
    ({"section": "startups", "start": "2026-10-09", "end": "2026-10-01"}, "after the To date"),
    ({"section": "linkedin", "source_id": 999}, "no longer a source"),
    ({"section": "linkedin", "stage": "seed"}, "by account"),
])
def test_bad_selections_are_refused_in_words(body, says) -> None:
    result = add(**body)
    assert not result["added"] and says in result["error"]


def test_the_same_selection_cannot_be_added_twice() -> None:
    assert add(section="startups", stage="seed")["added"]
    again = add(section="startups", stage="seed")
    assert not again["added"] and "already in Export" in again["error"]
    assert add(section="startups", stage="seed", start="2026-10-01")["added"]


def test_selections_can_be_taken_back_and_cleared() -> None:
    one = add(section="startups", stage="seed")["selection_id"]
    add(section="startups", stage="series-a")
    assert client.delete(f"/api/export/selections/{one}").json() == {"removed": True}
    assert client.delete(f"/api/export/selections/{one}").status_code == 404
    assert len(client.get("/api/export").json()["selections"]) == 1
    assert client.delete("/api/export/selections").json() == {"removed": 1}
    body = client.get("/api/export").json()
    assert body["selections"] == [] and body["rounds"] == [] and body["total"] == 0


def test_removing_a_linkedin_account_removes_its_selections() -> None:
    acme = linkedin_account()
    add(section="linkedin", source_id=acme)
    add(section="linkedin")
    li.remove_source(acme)
    assert [s["label"] for s in client.get("/api/export").json()["selections"]] == [
        "LinkedIn · all accounts"]


def test_excel_is_the_startup_firms_layout_with_every_source_named() -> None:
    from openpyxl import load_workbook

    feed_round("Orbotron", Stage.SEED, 1, round_label="Seed")
    acme = linkedin_account()
    linkedin_post(acme, LinkedinKind.FUNDING, "Orbotron", 2, round_label="Seed",
                  amount="$2M", investors="Blume Ventures")
    add(section="startups", stage="seed")
    add(section="linkedin")

    resp = client.get("/api/export.xlsx")
    assert resp.status_code == 200
    assert "gps-export-funding-" in resp.headers["content-disposition"]
    book = load_workbook(io.BytesIO(resp.content))
    assert book.sheetnames == list(STAGE_TITLES.values())
    for ws in book.worksheets:
        assert next(ws.iter_rows(values_only=True)) == tuple(name for name, _ in COLUMNS)
    rows = [r for r in book["Seed"].iter_rows(min_row=2, values_only=True) if r[0] == "Orbotron"]
    assert len(rows) == 1
    company, stage, round_, investors, amount, _, _, via, _, _ = rows[0]
    assert (stage, round_, investors, amount, via) == (
        "Seed", "Seed", "Blume Ventures", "$2M", "News feed, LinkedIn")
