"""Pieces shared by the platform's endpoints: date windows, paging, Excel."""

from __future__ import annotations

import io
from datetime import date

from fastapi import HTTPException
from fastapi.responses import StreamingResponse

from ..models import utcnow
from ..pipeline.timeutil import IST, day_bounds

MAX_LIMIT = 500
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def window(start: date | None, end: date | None):
    """[since, until) in UTC for inclusive Indian days. 400 when reversed."""
    if start and end and start > end:
        raise HTTPException(status_code=400, detail="The From date is after the To date.")
    return day_bounds(start, end)


def search_text(q: str | None) -> str | None:
    q = " ".join((q or "").split())
    return q[:200] or None


def sheet(workbook, title: str, columns: list[tuple[str, int]], rows: list[list],
          *, date_columns: dict[int, str], link_column: int | None) -> None:
    """One sheet with a frozen, filterable header.

    Every value is written as data, never as a formula: text scraped from the
    web that starts with "=" would otherwise run as one when the file opens.
    """
    from openpyxl.styles import Alignment, Font, PatternFill

    ws = workbook.create_sheet(title=title[:31])
    ws.append([name for name, _ in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1A00D9")
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"

    for values in rows:
        ws.append(values)
        n = ws.max_row
        for cell in ws[n]:
            if cell.data_type == "f":
                cell.data_type = "s"
        for column, fmt in date_columns.items():
            ws.cell(row=n, column=column).number_format = fmt
        if link_column is not None:
            link = ws.cell(row=n, column=link_column)
            if isinstance(link.value, str) and link.value.startswith(("http://", "https://")):
                link.hyperlink = link.value
                link.font = Font(color="1A00D9", underline="single")

    for index, (_, width) in enumerate(columns, start=1):
        ws.column_dimensions[ws.cell(row=1, column=index).column_letter].width = width
    ws.auto_filter.ref = ws.dimensions


def excel_date(value):
    """An aware datetime as a naive IST one: Excel has no time zones."""
    return value.astimezone(IST).replace(tzinfo=None) if value else None


def download(workbook, name: str) -> StreamingResponse:
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    day = utcnow().astimezone(IST).strftime("%Y-%m-%d")
    return StreamingResponse(buffer, media_type=XLSX, headers={
        "Content-Disposition": f'attachment; filename="gps-{name}-{day}.xlsx"'})
