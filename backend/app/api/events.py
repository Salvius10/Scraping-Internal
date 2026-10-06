"""Insights > Events organised: what the tracked VC firms run, on screen and as Excel.

  GET    /api/insights/events?when=upcoming|past   events, soonest first (upcoming)
                                                   or newest first (past)
  GET    /api/insights/events.xlsx                 the same rows as an Excel file
  POST   /api/insights/events/sources {url, label?, firm?}
                                                   add a website and read it now
  DELETE /api/insights/events/sources/{id}         remove a pasted website

Both GETs take `organiser` (a firm key, or a pasted source's key such as
"site-3") and `start` / `end` (YYYY-MM-DD, IST, inclusive) on the start date.

Rows come from `events`, which the scheduled refresh fills (`ingest/events.py`).
The GETs read the database only: no fetch, no model. Adding a source reads it
straight away, which is paid -- Apify for a Luma calendar, Firecrawl and
gpt-oss for any other page.
"""

from __future__ import annotations

import io
from datetime import date, datetime, timedelta

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..config import settings
from ..db import get_session
from ..ingest.events import (
    Feed, SourceError, add_source, all_feeds, read_now, remove_source,
)
from ..ingest.vc_firms import load_firms
from ..models import Event, EventRead, utcnow
from .insights import IST, Window

router = APIRouter(prefix="/api/insights", tags=["insights"])

WHEN = ("upcoming", "past")


class EventOut(BaseModel):
    id: int
    title: str
    url: str
    starts_at: datetime
    ends_at: datetime | None
    date_only: bool              # the source gave a day, no time
    timezone: str | None
    venue: str | None
    city: str | None
    country: str | None
    online: bool | None
    host: str | None
    description: str | None
    organiser: str               # the firm, or the pasted source's label
    organiser_key: str           # firm key, or source key for a source with no firm
    source_key: str
    via_label: str               # "Luma, via Apify" | "Website, via Firecrawl"


class SourceOut(BaseModel):
    key: str
    label: str
    url: str
    via: str                     # luma | page
    via_label: str
    firm: str | None
    custom: bool                 # pasted on the page, so it can be removed
    site: bool                   # the firm's own home page, checked weekly
    id: int | None
    india_only: bool
    note: str | None
    events: int
    last_read: datetime | None
    last_error: str | None


class OrganiserOut(BaseModel):
    key: str
    label: str
    count: int                   # events in the current tab and dates


class FirmRef(BaseModel):
    key: str
    label: str
    sources: int                 # event sources for this firm, registry and pasted
    no_events: str | None        # why it has none, when the registry knows


class EventsOut(BaseModel):
    when: str
    organiser: str | None
    counts: dict[str, int]       # {"upcoming": n, "past": n} for the current filters
    events: list[EventOut]
    organisers: list[OrganiserOut]
    sources: list[SourceOut]
    firms: list[FirmRef]         # every tracked firm, with or without a source
    apify_ready: bool            # APIFY_API_TOKEN is set
    firecrawl_ready: bool        # FIRECRAWL_API_KEY is set (or self-hosted)


def _organiser(feed: Feed) -> tuple[str, str]:
    """(key, label) an event is shown under: its firm, else the source itself."""
    if feed.firm:
        firm = next((f for f in load_firms() if f.key == feed.firm), None)
        return feed.firm, firm.label if firm else feed.label
    return feed.key, feed.label


def is_upcoming(event: Event, now: datetime) -> bool:
    """Still to come or under way. A day-only event lasts the whole IST day."""
    end = event.ends_at or (event.starts_at + timedelta(days=1) if event.date_only
                            else event.starts_at)
    return end >= now


def _rows(session: Session, window: Window, organiser: str | None
          ) -> tuple[dict[str, list[EventOut]], dict[str, dict[str, int]]]:
    """Events split into upcoming and past, and per-organiser counts for each."""
    feeds = {f.key: f for f in all_feeds()}
    now = utcnow()
    split: dict[str, list[EventOut]] = {w: [] for w in WHEN}
    per_org: dict[str, dict[str, int]] = {}
    seen: set[str] = set()

    # Registry sources first, so an event that is also in a pasted calendar
    # shows once, under the firm.
    rows = sorted(session.scalars(window.apply(select(Event), Event.starts_at)).all(),
                  key=lambda e: (e.source_key.startswith("site-"), e.id))
    for event in rows:
        feed = feeds.get(event.source_key)
        if feed is None:
            continue                     # a source removed from the registry
        link = event.url.rstrip("/").lower()
        if link != feed.url.rstrip("/").lower():   # the event's own page
            if link in seen:
                continue
            seen.add(link)

        org_key, org_label = _organiser(feed)
        when = "upcoming" if is_upcoming(event, now) else "past"
        per_org.setdefault(org_key, {w: 0 for w in WHEN})[when] += 1
        if organiser and org_key != organiser:
            continue
        split[when].append(EventOut(
            id=event.id, title=event.title, url=event.url, starts_at=event.starts_at,
            ends_at=event.ends_at, date_only=event.date_only, timezone=event.timezone,
            venue=event.venue, city=event.city, country=event.country, online=event.online,
            host=event.host, description=event.description, organiser=org_label,
            organiser_key=org_key, source_key=event.source_key,
            via_label=feed.via_label,
        ))

    split["upcoming"].sort(key=lambda e: e.starts_at)                 # soonest first
    split["past"].sort(key=lambda e: e.starts_at, reverse=True)       # newest first
    return split, per_org


def _sources(session: Session) -> list[SourceOut]:
    latest = (select(EventRead.source_key, func.max(EventRead.id).label("id"))
              .group_by(EventRead.source_key).subquery())
    last = {r.source_key: r for r in session.scalars(
        select(EventRead).join(latest, EventRead.id == latest.c.id)).all()}
    counts = dict(session.execute(
        select(Event.source_key, func.count(Event.id)).group_by(Event.source_key)).all())
    out = []
    for feed in all_feeds():
        read = last.get(feed.key)
        out.append(SourceOut(
            key=feed.key, label=_organiser(feed)[1], url=feed.url, via=feed.via,
            via_label=feed.via_label, firm=feed.firm, custom=feed.custom, site=feed.site,
            id=feed.source_id, india_only=feed.india_only, note=feed.note,
            events=counts.get(feed.key, 0),
            last_read=read.started_at if read else None,
            last_error=read.error if read else None,
        ))
    return out


def _organisers() -> dict[str, str]:
    """Every tracked firm, then each pasted source that is not tied to one."""
    out = {f.key: f.label for f in load_firms()}
    for feed in all_feeds():
        key, label = _organiser(feed)
        out.setdefault(key, label)
    return out


def _check(when: str, organiser: str | None) -> None:
    if when not in WHEN:
        raise HTTPException(status_code=400, detail="when is upcoming or past")
    if organiser and organiser not in _organisers():
        raise HTTPException(status_code=404, detail=f"unknown organiser {organiser!r}")


@router.get("/events", response_model=EventsOut)
def events(
    when: str = Query("upcoming", description="upcoming | past"),
    organiser: str | None = Query(None, description="a firm key, or a source key like site-3"),
    start: date | None = Query(None, description="starts on or after (YYYY-MM-DD, IST)"),
    end: date | None = Query(None, description="starts on or before (YYYY-MM-DD, IST)"),
    session: Session = Depends(get_session),
) -> EventsOut:
    _check(when, organiser)
    split, per_org = _rows(session, Window(start, end), organiser)

    if organiser:
        counts = dict(per_org.get(organiser, {w: 0 for w in WHEN}))
    else:
        counts = {w: sum(c[w] for c in per_org.values()) for w in WHEN}

    # Every tracked firm is listed, with or without a source or an event.
    feeds = all_feeds()
    sources_per_firm: dict[str, int] = {}
    for feed in feeds:
        if feed.firm:
            sources_per_firm[feed.firm] = sources_per_firm.get(feed.firm, 0) + 1

    return EventsOut(
        when=when, organiser=organiser, counts=counts, events=split[when],
        organisers=sorted(
            (OrganiserOut(key=k, label=v, count=per_org.get(k, {}).get(when, 0))
             for k, v in _organisers().items()),
            key=lambda o: o.label.lower()),
        sources=_sources(session),
        firms=sorted((FirmRef(key=f.key, label=f.label,
                              sources=sources_per_firm.get(f.key, 0), no_events=f.no_events)
                      for f in load_firms()),
                     key=lambda f: f.label.lower()),
        apify_ready=bool(settings.apify_api_token),
        firecrawl_ready=bool(settings.firecrawl_api_key)
        or "api.firecrawl.dev" not in settings.firecrawl_api_url,
    )


# --- Excel ------------------------------------------------------------------------

COLUMNS = [
    ("Organiser", 26), ("Event", 56), ("Starts (IST)", 19), ("Ends (IST)", 19),
    ("City", 16), ("Venue", 30), ("Format", 10), ("Found via", 22), ("Link", 50),
]


def _ist(value: datetime | None):
    return value.astimezone(IST).replace(tzinfo=None) if value else None


@router.get("/events.xlsx")
def events_excel(
    when: str = Query("upcoming"),
    organiser: str | None = Query(None),
    start: date | None = Query(None),
    end: date | None = Query(None),
    session: Session = Depends(get_session),
) -> StreamingResponse:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    _check(when, organiser)
    window = Window(start, end)
    split, _ = _rows(session, window, organiser)

    workbook = Workbook()
    ws = workbook.active
    ws.title = "Upcoming events" if when == "upcoming" else "Past events"
    ws.append([name for name, _ in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1A00D9")
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"

    for e in split[when]:
        fmt = "Online" if e.online else "In person" if e.online is False else ""
        ws.append([e.organiser, e.title, _ist(e.starts_at), _ist(e.ends_at), e.city or "",
                   e.venue or "", fmt, e.via_label, e.url])
        row = ws.max_row
        for column in (3, 4):
            ws.cell(row=row, column=column).number_format = (
                "dd mmm yyyy" if e.date_only else "dd mmm yyyy hh:mm")
        link = ws.cell(row=row, column=9)
        link.hyperlink = e.url
        link.font = Font(color="1A00D9", underline="single")

    for index, (_, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=index).column_letter].width = width
    ws.auto_filter.ref = ws.dimensions

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    day = utcnow().astimezone(IST).strftime("%Y-%m-%d")
    name = f"gps-events-{when}-{organiser or 'all'}-{window.label()}-{day}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )


# --- Pasted sources ---------------------------------------------------------------

class AddSourceIn(BaseModel):
    url: str
    label: str | None = None
    firm: str | None = None


class AddSourceOut(BaseModel):
    added: bool
    source_key: str | None = None
    via: str | None = None
    found: int = 0               # events on the source
    new: int = 0                 # of which not seen before
    apify_usd: float = 0.0
    credits_used: int = 0
    llm_usd: float = 0.0
    error: str | None = None


@router.post("/events/sources", response_model=AddSourceOut)
def add_events_source(request: AddSourceIn) -> AddSourceOut:
    """Save a website as a source, then read it straight away (paid)."""
    try:
        feed = add_source(request.url, request.label, request.firm)
    except SourceError as exc:
        return AddSourceOut(added=False, error=str(exc))

    result = read_now(feed)
    return AddSourceOut(
        added=True, source_key=feed.key, via=feed.via,
        found=sum(r.items_seen for r in result.reads), new=result.new,
        apify_usd=round(result.apify_usd, 6), credits_used=result.credits,
        llm_usd=round(result.llm_usd, 6),
        error=next((r.error for r in result.reads if r.error), None),
    )


@router.delete("/events/sources/{source_id}")
def delete_events_source(source_id: int) -> dict:
    if not remove_source(source_id):
        raise HTTPException(status_code=404, detail="no such source")
    return {"removed": True}
