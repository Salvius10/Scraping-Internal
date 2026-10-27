"""Platform: Export to Excel.

  GET /api/export/funding.xlsx?stage=&source=&start=&end=&q=
  GET /api/export/events.xlsx?when=&city=&start=&end=&q=
  GET /api/export/enrichment.xlsx?designation=&q=

Each takes the same filters as its feed, so the file holds exactly what the
reader is looking at -- every page of it. Every row carries its source URL
(clickable) and date. Reads only: no fetch, no model, so it costs nothing.

The enrichment file accounts for every company sent: its leaders first, then
a row for each company with none to show, and why (no leaders found, failed,
waiting, or none with the chosen designation).
"""

from __future__ import annotations

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from ..db import get_session
from ..enrichment import roles
from ..models import EnrichStatus
from . import enrichment as enrichment_api
from . import events as events_api
from . import funding as funding_api
from .common import download, excel_date, sheet

router = APIRouter(prefix="/api/export", tags=["export"])

EXPORT_LIMIT = 50_000

FUNDING_COLUMNS = [
    ("Company", 28), ("Stage", 11), ("Round", 16), ("Amount", 18), ("Investors", 40),
    ("Sector", 18), ("HQ city", 14), ("Summary", 60), ("Published (IST)", 18),
    ("Source", 22), ("Found via", 14), ("Source URL", 60), ("Also reported at", 60),
]

EVENT_COLUMNS = [
    ("Event", 44), ("Starts (IST)", 18), ("Ends (IST)", 18), ("City", 14), ("Venue", 30),
    ("Online", 8), ("Organiser", 28), ("Category", 16), ("Price", 14), ("Summary", 60),
    ("Source", 14), ("Event URL", 60), ("Also listed at", 60),
]

ENRICHMENT_COLUMNS = [
    ("Name", 26), ("Designation", 22), ("Company", 26), ("Email", 32), ("Phone", 18),
    ("LinkedIn", 48), ("Status", 18), ("Title", 36), ("Other emails", 32),
    ("Other phones", 22), ("Location", 26), ("Company LinkedIn page", 40),
    ("Company website", 28), ("Check", 48), ("Funding stage", 13), ("Sector", 18),
    ("HQ city", 14), ("Enriched (IST)", 18),
]
MAYBE_OTHER = "May be at another company with the same name"

# The Status column, for a company with no leader rows in the file.
NOT_ENRICHED = {
    EnrichStatus.NO_MATCH: "No leaders found",
    EnrichStatus.FAILED: "Failed",
    EnrichStatus.QUEUED: "Waiting",
    EnrichStatus.WORKING: "Waiting",
}


@router.get("/funding.xlsx")
def funding_excel(f: funding_api.FundingFilter = Depends(funding_api.funding_filter),
                  session: Session = Depends(get_session)) -> StreamingResponse:
    from openpyxl import Workbook

    rows = session.scalars(funding_api.ordered(funding_api.query(f)).limit(EXPORT_LIMIT)).all()
    values = [[
        r.company, r.stage.value, r.round_label or "", r.amount or "Undisclosed",
        ", ".join(r.investors or []), r.sector or "", r.hq_city or "", r.summary or "",
        excel_date(r.published_at), r.source_label,
        funding_api.SOURCE_TYPES.get(r.source_type, r.source_type), r.source_url,
        "\n".join(s["url"] for s in (r.sources or [])
                  if isinstance(s, dict) and s.get("url") and s["url"] != r.source_url),
    ] for r in rows]
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheet(workbook, "Funded startups", FUNDING_COLUMNS, values,
          date_columns={9: "dd mmm yyyy hh:mm"}, link_column=12)
    return download(workbook, f"funded-startups-{f.stage or 'all-stages'}")


@router.get("/events.xlsx")
def events_excel(f: events_api.EventFilter = Depends(events_api.event_filter),
                 session: Session = Depends(get_session)) -> StreamingResponse:
    from openpyxl import Workbook

    rows = session.scalars(events_api.ordered(events_api.query(f), f.when)
                           .limit(EXPORT_LIMIT)).all()
    values = [[
        r.title, excel_date(r.starts_at), excel_date(r.ends_at), r.city or "", r.venue or "",
        "Yes" if r.online else ("No" if r.online is False else ""), r.organiser or "",
        r.category or "", r.price or "", r.summary or "", r.source_label, r.url,
        "\n".join(s["url"] for s in (r.sources or [])
                  if isinstance(s, dict) and s.get("url") and s["url"] != r.url),
    ] for r in rows]
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheet(workbook, "India startup events", EVENT_COLUMNS, values,
          date_columns={2: "dd mmm yyyy hh:mm", 3: "dd mmm yyyy hh:mm"}, link_column=12)
    return download(workbook, f"startup-events-{f.when}")


@router.get("/enrichment.xlsx")
def enrichment_excel(f: enrichment_api.ContactFilter = Depends(enrichment_api.contact_filter),
                     session: Session = Depends(get_session)) -> StreamingResponse:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    rows = session.execute(enrichment_api.ordered(enrichment_api.query(f))
                           .limit(EXPORT_LIMIT)).all()
    values = [[
        c.full_name, c.designation, company.company, c.email or "", c.phone or "",
        c.linkedin or "",
        "Enriched" if company.status == EnrichStatus.DONE else "Partly enriched",
        c.title or "", "\n".join(c.other_emails or []), "\n".join(c.other_phones or []),
        c.location or "", c.company_page or "", c.company_site or "",
        MAYBE_OTHER if c.maybe_other else "", company.stage or "", company.sector or "",
        company.hq_city or "", excel_date(c.looked_up_at),
    ] for c, company in rows]
    leaders = len(values)

    wanted = roles.DESIGNATION_KEYS.get(f.designation or "")
    for company in enrichment_api.not_shown(session, f, {company.id for _, company in rows}):
        if company.status == EnrichStatus.DONE and wanted:  # enriched; none with that designation
            status, why = f"No {wanted} found", f"Enriched, but SignalHire found no {wanted} here."
        else:
            status = NOT_ENRICHED.get(company.status, "Enriched")
            why = company.note or "Waiting to be enriched."
        values.append(["", "", company.company, "", "", "", status, "", "", "", "", "", "", why,
                       company.stage or "", company.sector or "", company.hq_city or "",
                       excel_date(company.enriched_at)])

    workbook = Workbook()
    workbook.remove(workbook.active)
    sheet(workbook, "Enriched contacts", ENRICHMENT_COLUMNS, values,
          date_columns={18: "dd mmm yyyy hh:mm"}, link_column=6)
    # The companies with no leaders to show, set apart in grey below the leaders.
    ws = workbook["Enriched contacts"]
    for row in ws.iter_rows(min_row=leaders + 2):
        for cell in row:
            cell.font = Font(color="656C81", italic=True)
    return download(workbook, f"enriched-contacts-{f.designation or 'all-designations'}")
