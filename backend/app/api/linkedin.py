"""LinkedIn: funding and news posted by the accounts the reader adds.

  GET    /api/linkedin?source=3&kind=funding    posts newest first, every account
                                                with its count and last read
  GET    /api/linkedin.xlsx?...                 the same rows as an Excel file
  POST   /api/linkedin/sources {url, label?}    add a profile or company page,
                                                then read it now
  DELETE /api/linkedin/sources/{id}             remove it, with what was read

The GETs take `start` / `end` (YYYY-MM-DD, IST, inclusive) on the post time,
like the Insights pages, `kind` (funding | news) to show one kind, and
`all=true` to include posts read as Other. They read only: no fetch, no model.

Adding reads the account straight away through Apify (~$0.04 for the newest
20 posts), then gpt-oss reads them (~$0.00007 a post). See ingest/linkedin.py.
"""

from __future__ import annotations

import io
from datetime import date, datetime

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..db import get_session, session_scope
from ..ingest.linkedin import (
    KIND_LABELS, SourceError, add_source, all_sources, read_now, ready, remove_source,
)
from ..models import LinkedinKind, LinkedinPost, LinkedinRead, utcnow
from .insights import IST, Window

router = APIRouter(prefix="/api/linkedin", tags=["linkedin"])

KINDS = {"funding": LinkedinKind.FUNDING, "news": LinkedinKind.NEWS}
SNIPPET_CHARS = 400


class PostOut(BaseModel):
    id: int
    source_id: int
    account: str
    author: str | None           # who wrote it, when the account reposted it
    repost: bool
    kind: str                    # Funding | News | Other
    headline: str
    snippet: str                 # the post's own words, cut short
    company: str | None
    round: str | None
    amount: str | None
    investors: list[str]
    posted_at: datetime | None
    url: str


class SourceOut(BaseModel):
    id: int
    label: str
    url: str
    kind: str                    # profile | company
    kind_label: str
    count: int                   # posts shown for the current filters
    posts: int                   # every post read from it
    last_read: datetime | None
    last_error: str | None


class LinkedinOut(BaseModel):
    source: int | None
    kind: str | None
    counts: dict[str, int]       # funding, news, other -- for the dates and account
    posts: list[PostOut]
    sources: list[SourceOut]
    pending: int                 # posts read but not yet sorted by the model
    hidden_other: int            # posts read as Other, not shown
    apify_ready: bool


def _kind(key: str | None) -> LinkedinKind | None:
    if not key:
        return None
    try:
        return KINDS[key]
    except KeyError:
        raise HTTPException(status_code=400, detail="kind is funding or news") from None


def _labels() -> dict[int, str]:
    return {s.id: s.label for s in all_sources()}


def _check_source(source: int | None) -> int | None:
    if source is not None and source not in _labels():
        raise HTTPException(status_code=404, detail=f"unknown LinkedIn source {source}")
    return source


def _snippet(text: str) -> str:
    flat = " ".join(text.split())
    return flat if len(flat) <= SNIPPET_CHARS else flat[:SNIPPET_CHARS].rsplit(" ", 1)[0] + "…"


def _posts(session: Session, window: Window) -> list[PostOut]:
    """Every sorted post in the date window, newest first; undated ones last."""
    labels = _labels()
    stmt = window.apply(
        select(LinkedinPost).where(LinkedinPost.classified_at.is_not(None),
                                   LinkedinPost.source_id.in_(list(labels))),
        LinkedinPost.posted_at)
    stmt = stmt.order_by(LinkedinPost.posted_at.desc().nullslast(), LinkedinPost.id.desc())
    return [
        PostOut(
            id=p.id, source_id=p.source_id, account=labels[p.source_id],
            author=p.author if p.repost else None, repost=p.repost,
            kind=(p.kind or LinkedinKind.OTHER).value,
            headline=p.headline or _snippet(p.content)[:160], snippet=_snippet(p.content),
            company=p.company, round=p.round_label, amount=p.amount,
            investors=[n for n in (p.investors or "").split("; ") if n],
            posted_at=p.posted_at, url=p.url,
        )
        for p in session.scalars(stmt).all()
    ]


def _rows(session: Session, window: Window, source: int | None, kind: LinkedinKind | None,
          include_other: bool) -> tuple[list[PostOut], dict[str, int], dict[int, int]]:
    """(rows to show, counts per kind for the account, rows per account for the kind)."""
    posts = _posts(session, window)
    wanted = ({kind.value} if kind else
              {LinkedinKind.FUNDING.value, LinkedinKind.NEWS.value}
              | ({LinkedinKind.OTHER.value} if include_other else set()))
    counts = {k.value.lower(): 0 for k in LinkedinKind}
    per_source: dict[int, int] = {}
    rows = []
    for p in posts:
        if source is None or p.source_id == source:
            counts[p.kind.lower()] += 1
        if p.kind not in wanted:
            continue
        per_source[p.source_id] = per_source.get(p.source_id, 0) + 1
        if source is None or p.source_id == source:
            rows.append(p)
    return rows, counts, per_source


def sources_out(session: Session, per_source: dict[int, int]) -> list[SourceOut]:
    sources = all_sources()
    if not sources:
        return []
    latest = (select(LinkedinRead.source_id, func.max(LinkedinRead.id).label("id"))
              .group_by(LinkedinRead.source_id).subquery())
    last = {r.source_id: r for r in session.scalars(
        select(LinkedinRead).join(latest, LinkedinRead.id == latest.c.id)).all()}
    totals = dict(session.execute(
        select(LinkedinPost.source_id, func.count(LinkedinPost.id))
        .group_by(LinkedinPost.source_id)).all())
    out = []
    for src in sources:
        read = last.get(src.id)
        out.append(SourceOut(
            id=src.id, label=src.label, url=src.url, kind=src.kind,
            kind_label=KIND_LABELS.get(src.kind, src.kind), count=per_source.get(src.id, 0),
            posts=totals.get(src.id, 0), last_read=read.started_at if read else None,
            last_error=read.error if read else None,
        ))
    return out


@router.get("", response_model=LinkedinOut)
def linkedin(
    source: int | None = Query(None, description="one account's id, or omit for all"),
    kind: str | None = Query(None, description="funding | news, or omit for both"),
    start: date | None = Query(None, description="posted on or after (YYYY-MM-DD, IST)"),
    end: date | None = Query(None, description="posted on or before (YYYY-MM-DD, IST)"),
    all: bool = Query(False, description="include posts read as Other"),
    session: Session = Depends(get_session),
) -> LinkedinOut:
    source = _check_source(source)
    chosen = _kind(kind)
    rows, counts, per_source = _rows(session, Window(start, end), source, chosen, all)
    pending = session.scalar(
        select(func.count(LinkedinPost.id)).where(LinkedinPost.classified_at.is_(None))) or 0
    return LinkedinOut(
        source=source, kind=kind or None, counts=counts, posts=rows,
        sources=sources_out(session, per_source), pending=pending,
        hidden_other=0 if all or chosen else counts["other"], apify_ready=ready(),
    )


# --- Adding and removing ----------------------------------------------------------------

class AddSourceIn(BaseModel):
    url: str
    label: str | None = None


class AddSourceOut(BaseModel):
    added: bool
    source_id: int | None = None
    label: str | None = None
    found: int = 0               # posts read
    new: int = 0                 # of which not seen before
    funding: int = 0             # of the new posts, read as Funding
    news: int = 0                # ... and as News
    apify_usd: float = 0.0       # estimated
    llm_usd: float = 0.0
    error: str | None = None


@router.post("/sources", response_model=AddSourceOut)
def add_linkedin_source(request: AddSourceIn) -> AddSourceOut:
    """Save a LinkedIn account as a source, then read it straight away."""
    try:
        source = add_source(request.url, request.label)
    except SourceError as exc:
        return AddSourceOut(added=False, error=str(exc))

    result = read_now(source)
    read = result.reads[0]
    with session_scope() as s:
        kinds = dict(s.execute(
            select(LinkedinPost.kind, func.count(LinkedinPost.id))
            .where(LinkedinPost.source_id == source.id, LinkedinPost.kind.is_not(None))
            .group_by(LinkedinPost.kind)).all())
    stopped = result.classify.stopped_reason if result.classify else None
    return AddSourceOut(
        added=True, source_id=source.id, label=source.label, found=read.items_seen,
        new=read.items_new, funding=kinds.get(LinkedinKind.FUNDING, 0),
        news=kinds.get(LinkedinKind.NEWS, 0), apify_usd=round(read.apify_usd, 6),
        llm_usd=round(result.llm_usd, 6),
        error=read.error or (f"The posts were not read yet: {stopped}" if stopped else None),
    )


@router.delete("/sources/{source_id}")
def delete_linkedin_source(source_id: int) -> dict:
    if not remove_source(source_id):
        raise HTTPException(status_code=404, detail="no such source")
    return {"removed": True}


# --- Excel ------------------------------------------------------------------------

COLUMNS = [
    ("Account", 24), ("What", 10), ("Company", 26), ("Round", 16), ("Investors", 36),
    ("Amount", 18), ("Posted (IST)", 18), ("Headline", 50), ("Post", 70), ("Link", 50),
]


@router.get(".xlsx")
def linkedin_excel(
    source: int | None = Query(None),
    kind: str | None = Query(None),
    start: date | None = Query(None),
    end: date | None = Query(None),
    all: bool = Query(False),
    session: Session = Depends(get_session),
) -> StreamingResponse:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    source = _check_source(source)
    window = Window(start, end)
    rows, _, _ = _rows(session, window, source, _kind(kind), all)

    workbook = Workbook()
    ws = workbook.active
    ws.title = (_labels()[source] if source is not None else "LinkedIn")[:31]
    ws.append([name for name, _ in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1A00D9")
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"

    for r in rows:
        posted = r.posted_at.astimezone(IST).replace(tzinfo=None) if r.posted_at else None
        account = f"{r.account} (repost of {r.author})" if r.repost and r.author else r.account
        ws.append([account, r.kind, r.company or "", r.round or "", ", ".join(r.investors),
                   r.amount or "", posted, r.headline, r.snippet, r.url])
        row = ws.max_row
        ws.cell(row=row, column=7).number_format = "dd mmm yyyy hh:mm"
        link = ws.cell(row=row, column=10)
        link.hyperlink = r.url
        link.font = Font(color="1A00D9", underline="single")

    for index, (_, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=index).column_letter].width = width
    ws.auto_filter.ref = ws.dimensions

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    day = utcnow().astimezone(IST).strftime("%Y-%m-%d")
    name = f"gps-linkedin-{kind or 'all'}-{window.label()}-{day}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )
